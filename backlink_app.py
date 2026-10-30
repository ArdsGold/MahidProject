"""Native Windows UI for the deterministic backlink publisher; no model or AI service."""
import csv
import json
import os
import queue
import shutil
import subprocess
import sys
import threading
from pathlib import Path
import tkinter as tk
from tkinter import ttk, filedialog, messagebox

ROOT = Path(__file__).resolve().parent
DEFAULT_WORKBOOK = ROOT / '.local/workbooks/Cheer_Athletics_St_Louis.xlsx'
DEFAULT_POOL = ROOT / 'websites/Website_Pool.xlsx'
LABELS = {'published': 'Verified', 'needs_mapping': 'Needs mapping', 'not_run': 'Not run',
          'uncertain': 'Needs inspection', 'submitting': 'Submission in progress',
          'unverified': 'Needs inspection', 'needs_attention': 'Needs attention'}


def inspect_client(workbook, pool):
    from backlinks import load_workbook, campaign_folder
    from website_pool import merge_pool
    from mappings.registry import SUPPORTED
    business, entries = load_workbook(Path(workbook))
    entries = merge_pool(entries, Path(pool), Path(workbook))
    return {'name': business['name'], 'folder': campaign_folder(business),
            'hosts': list(dict.fromkeys(e['domain'] for e in entries)),
            'ready': len({e['domain'] for e in entries} & SUPPORTED)}


def read_state(folder):
    path = Path(folder) / 'state.json'
    return json.loads(path.read_text(encoding='utf-8')) if path.exists() else {}


def latest_report(folder):
    files = list(Path(folder).glob('results*.csv'))
    if not files:
        raise FileNotFoundError('No results report exists yet. Start a run first.')
    return max(files, key=lambda p: p.stat().st_mtime_ns)


def export_results(folder, destination):
    """Export public reports only; private account/checkpoint files are never included."""
    folder, destination = Path(folder), Path(destination)
    if destination.suffix.lower() not in ('.xlsx', '.csv'):
        raise ValueError('Choose an XLSX or CSV destination')
    report = latest_report(folder)
    if destination.suffix.lower() == '.csv':
        if report.resolve() != destination.resolve():
            shutil.copyfile(report, destination)
        return
    from openpyxl import Workbook
    workbook = Workbook()
    results = workbook.active
    results.title = 'Results'
    with report.open(encoding='utf-8-sig', newline='') as source:
        for row in csv.reader(source):
            results.append(["'" + value if value.startswith(('=', '+', '-', '@')) else value for value in row])
    verified = workbook.create_sheet('Verified Links')
    verified.append(['Verified public backlink URL'])
    urls = folder / 'published_urls.txt'
    if urls.exists():
        for url in dict.fromkeys(urls.read_text(encoding='utf-8').splitlines()):
            if url:
                verified.append([url])
    for sheet in workbook:
        sheet.freeze_panes = 'A2'
        sheet.auto_filter.ref = sheet.dimensions
        for column in sheet.columns:
            sheet.column_dimensions[column[0].column_letter].width = min(70, max(18, max(len(str(c.value or '')) for c in column) + 2))
    workbook.save(destination)
    workbook.close()


class Worker:
    def __init__(self, arguments, events):
        executable = Path(sys.executable)
        if executable.name.lower() == 'pythonw.exe':
            executable = executable.with_name('python.exe')
        self.process = subprocess.Popen([str(executable), '-u', str(ROOT / 'desktop_runner.py'), *arguments],
            cwd=ROOT, stdin=subprocess.PIPE, stdout=subprocess.PIPE, stderr=subprocess.DEVNULL,
            text=True, encoding='utf-8', creationflags=getattr(subprocess, 'CREATE_NO_WINDOW', 0))
        threading.Thread(target=self._read, args=(events,), daemon=True).start()

    def _read(self, events):
        try:
            for line in self.process.stdout:
                try:
                    value = json.loads(line)
                    if isinstance(value, dict) and 'event' in value:
                        events.put(value)
                except ValueError:
                    pass
        finally:
            code = self.process.wait()
            self.process.stdout.close()
            self.process.stdin.close()
            events.put({'event': 'exit', 'code': code})

    def send(self, command, **values):
        if self.process.poll() is None:
            self.process.stdin.write(json.dumps({'command': command, **values}) + '\n')
            self.process.stdin.flush()


class BacklinkApp:
    def __init__(self, root):
        self.root, self.worker, self.client, self.action_id = root, None, None, None
        self.events, self.closing, self.diagnostic = queue.Queue(), False, None
        self.validation_error = 'Choose a valid completed client XLSX and website pool.'
        root.title('Backlink Publisher')
        root.geometry('1120x760')
        root.minsize(880, 620)
        style = ttk.Style(root)
        if 'vista' in style.theme_names():
            style.theme_use('vista')
        style.configure('Title.TLabel', font=('Segoe UI', 22, 'bold'))
        style.configure('Status.TLabel', font=('Segoe UI', 11))
        style.configure('Treeview', rowheight=28)
        frame = ttk.Frame(root, padding=22)
        frame.pack(fill='both', expand=True)
        ttk.Label(frame, text='Backlink Publisher', style='Title.TLabel').pack(anchor='w')
        ttk.Label(frame, text='Saved workflows. Local automation. No AI service.').pack(anchor='w', pady=(2, 18))
        self.workbook = tk.StringVar(value=str(DEFAULT_WORKBOOK))
        self.pool = tk.StringVar(value=str(DEFAULT_POOL))
        self.file_buttons = []
        for label, variable in (('Client workbook', self.workbook), ('Website pool', self.pool)):
            row = ttk.Frame(frame)
            row.pack(fill='x', pady=4)
            ttk.Label(row, text=label, width=18).pack(side='left')
            ttk.Entry(row, textvariable=variable, state='readonly').pack(side='left', fill='x', expand=True)
            button = ttk.Button(row, text='Choose XLSX', command=lambda v=variable: self.choose(v))
            button.pack(side='left', padx=(8, 0))
            self.file_buttons.append(button)
        self.summary = tk.StringVar(value='Choose a completed client workbook to begin.')
        ttk.Label(frame, textvariable=self.summary, style='Status.TLabel').pack(anchor='w', pady=(16, 5))
        self.progress_text = tk.StringVar(value='0 / 100 verified websites')
        ttk.Label(frame, textvariable=self.progress_text).pack(anchor='w')
        self.progress = ttk.Progressbar(frame, maximum=100)
        self.progress.pack(fill='x', pady=(5, 12))
        buttons = ttk.Frame(frame)
        buttons.pack(fill='x')
        self.connect_button = ttk.Button(buttons, text='Connect Chrome', command=lambda: self.launch(connect=True))
        self.connect_button.pack(side='left')
        self.start_button = ttk.Button(buttons, text='Start / Resume', command=self.launch)
        self.start_button.pack(side='left', padx=8)
        self.pause_button = ttk.Button(buttons, text='Pause', state='disabled', command=self.pause)
        self.pause_button.pack(side='left')
        ttk.Button(buttons, text='Export results', command=self.export).pack(side='right')
        ttk.Button(buttons, text='Open results folder', command=self.open_results).pack(side='right', padx=8)
        self.status = tk.StringVar(value='Ready. Keep Ui.Vision open in regular Chrome. Runs continue through the full pool.')
        ttk.Label(frame, textvariable=self.status, wraplength=1000).pack(anchor='w', pady=(12, 8))
        self.action_frame = ttk.LabelFrame(frame, text='Your action is needed', padding=12)
        self.action_message = tk.StringVar()
        ttk.Label(self.action_frame, textvariable=self.action_message, wraplength=980).pack(anchor='w')
        self.public_url = tk.StringVar()
        self.url_row = ttk.Frame(self.action_frame)
        ttk.Label(self.url_row, text='Public listing URL').pack(side='left')
        ttk.Entry(self.url_row, textvariable=self.public_url).pack(side='left', fill='x', expand=True, padx=8)
        ttk.Button(self.url_row, text='Verify URL', command=lambda: self.respond(self.public_url.get().strip())).pack(side='left')
        action_buttons = ttk.Frame(self.action_frame)
        action_buttons.pack(fill='x', pady=(10, 0))
        ttk.Button(action_buttons, text='Done / Continue', command=lambda: self.respond('')).pack(side='left')
        ttk.Button(action_buttons, text='Skip this site', command=lambda: self.respond('skip')).pack(side='left', padx=8)
        ttk.Label(action_buttons, text='Complete verification directly in Chrome; leave submission to the app.').pack(side='left', padx=8)
        table = ttk.Frame(frame)
        table.pack(fill='both', expand=True, pady=(12, 0))
        self.table = ttk.Treeview(table, columns=('site', 'status', 'url', 'reason'), show='headings')
        for column, title, width in (('site', 'Website', 210), ('status', 'Result', 135), ('url', 'Verified public URL', 310), ('reason', 'Notes', 340)):
            self.table.heading(column, text=title)
            self.table.column(column, width=width, minwidth=90)
        scrollbar = ttk.Scrollbar(table, orient='vertical', command=self.table.yview)
        self.table.configure(yscrollcommand=scrollbar.set)
        scrollbar.pack(side='right', fill='y')
        self.table.pack(fill='both', expand=True)
        root.protocol('WM_DELETE_WINDOW', self.close)
        self.last_state = None
        self.refresh_client()
        root.after(250, self.tick)

    def choose(self, variable):
        path = filedialog.askopenfilename(parent=self.root, filetypes=[('Excel workbooks', '*.xlsx')])
        if path:
            variable.set(path)
            self.refresh_client()

    def refresh_client(self):
        try:
            self.client = inspect_client(self.workbook.get(), self.pool.get())
            total = len(self.client['hosts'])
            self.summary.set(f"{self.client['name']}  |  {total} enabled websites  |  {self.client['ready']} saved workflows  |  {total - self.client['ready']} need mapping")
            self.last_state = None
            self.refresh_results()
        except Exception as error:
            self.client = None
            self.validation_error = str(error)[:400] if isinstance(error, (ValueError, OSError)) else 'The workbook could not be read. Check its Business, Hours and Sites sheets.'
            self.summary.set('Select a valid completed client workbook and website pool.')
            self.table.delete(*self.table.get_children())
            self.progress['value'] = 0
            self.progress_text.set('No client loaded')

    def refresh_results(self):
        if not self.client:
            return
        from website_pool import verified_count
        from mappings.registry import SUPPORTED
        state = read_state(self.client['folder'])
        if state == self.last_state:
            return
        self.last_state = state
        count = verified_count(state)
        self.progress['value'] = min(count, 100)
        self.progress_text.set(f'{count} / 100 verified websites · full pool mode')
        self.table.delete(*self.table.get_children())
        for host in dict.fromkeys([*self.client['hosts'], *state]):
            record = state.get(host, {})
            status = record.get('status', 'not_run')
            label = LABELS.get(status, status.replace('_', ' ').capitalize())
            if host not in SUPPORTED and status in ('not_run', 'needs_mapping', 'needs_attention'):
                label = 'Needs mapping'
            self.table.insert('', 'end', values=(host, label, record.get('url', '') if status == 'published' else '', record.get('reason', '')))

    def set_running(self, running):
        for widget in [self.start_button, self.connect_button, *self.file_buttons]:
            widget.configure(state='disabled' if running else 'normal')
        self.pause_button.configure(state='normal' if running else 'disabled')

    def launch(self, connect=False):
        if self.worker:
            return
        if not connect:
            self.refresh_client()
            if not self.client:
                messagebox.showerror('Workbook needs attention', self.validation_error, parent=self.root)
                return
        arguments = ['--connect'] if connect else ['--workbook', self.workbook.get(), '--pool', self.pool.get()]
        try:
            self.worker = Worker(arguments, self.events)
            self.set_running(True)
            self.status.set('Connecting Chrome…' if connect else 'Starting or resuming the saved campaign…')
        except OSError:
            self.worker = None
            self.status.set('Python could not start the worker. Check the local Python installation.')

    def pause(self):
        if self.worker:
            try:
                self.worker.send('stop')
            except (OSError, ValueError):
                pass
            self.pause_button.configure(state='disabled')
            self.status.set('Pausing at the next safe checkpoint. Chrome stays open and progress is preserved.')

    def respond(self, value):
        if self.worker and self.action_id:
            try:
                self.worker.send('respond', id=self.action_id, value=value)
                self.action_id = None
                self.action_frame.pack_forget()
                self.status.set('Continuing…')
            except (OSError, ValueError):
                self.status.set('The worker has stopped. Your saved progress is preserved.')

    def tick(self):
        while not self.events.empty():
            event = self.events.get()
            kind = event['event']
            if kind == 'action':
                self.action_id = event['id']
                self.action_message.set(event['message'])
                self.public_url.set('')
                if event.get('kind') == 'review':
                    self.url_row.pack(fill='x', pady=(8, 0))
                else:
                    self.url_row.pack_forget()
                self.action_frame.pack(fill='x', before=self.table.master, pady=(4, 0))
                self.status.set('Waiting for you. Complete the action in the browser, then click Done / Continue.')
            elif kind == 'site':
                self.status.set('Working on ' + event['host'])
            elif kind == 'diagnostic':
                self.diagnostic = event['path']
            elif kind in ('connected', 'finished', 'stopped', 'error', 'review'):
                self.status.set(event['message'])
            elif kind == 'exit':
                self.worker = None
                self.action_id = None
                self.action_frame.pack_forget()
                self.set_running(False)
                if event['code']:
                    self.status.set('The worker ended unexpectedly. Progress is saved; review results before resuming.')
                if self.closing:
                    self.root.destroy()
                    return
        try:
            self.refresh_results()
        except (OSError, ValueError):
            pass  # Keep the last complete view while another process saves atomically.
        self.root.after(500, self.tick)

    def open_results(self):
        if self.client:
            folder = self.client['folder']
            folder.mkdir(parents=True, exist_ok=True)
            os.startfile(folder)

    def export(self):
        if not self.client:
            return
        destination = filedialog.asksaveasfilename(parent=self.root, defaultextension='.xlsx',
            initialfile=self.client['folder'].name + '-results.xlsx', filetypes=[('Excel workbook', '*.xlsx'), ('CSV report', '*.csv')])
        if destination:
            try:
                if Path(destination).resolve() in {Path(self.workbook.get()).resolve(), Path(self.pool.get()).resolve()}:
                    messagebox.showerror('Choose another destination', 'Export cannot overwrite the client workbook or website pool.', parent=self.root)
                    return
                export_results(self.client['folder'], destination)
                self.status.set('Results exported. Verified Links contains verified public URLs only.')
            except (OSError, ValueError):
                messagebox.showerror('Export unavailable', 'No report is ready, or the destination is locked. Close Excel and try again.', parent=self.root)

    def close(self):
        if self.worker:
            self.closing = True
            self.pause()
        else:
            self.root.destroy()


def main():
    root = tk.Tk()
    BacklinkApp(root)
    root.mainloop()


if __name__ == '__main__':
    main()
