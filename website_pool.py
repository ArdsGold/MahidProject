"""Editable website pool, per-campaign attempt history and a distinct-domain target."""
import argparse
import csv
import json
import sqlite3
import uuid
from collections import Counter
from contextlib import contextmanager
from datetime import datetime, timezone
from pathlib import Path

import openpyxl
from openpyxl.styles import Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter

from mappings.common import domain
from mappings.registry import SUPPORTED, DIRECTORIES, CLASSIFIEDS

ROOT = Path(__file__).resolve().parent
DEFAULT_POOL = ROOT / 'websites' / 'Website_Pool.xlsx'
DEFAULT_TARGET = 100
HEADERS = ('Website', 'Enabled', 'Notes', 'Domain', 'Workflow', 'Confidence %',
           'Successes', 'Errors', 'Recovered links', 'Deferred / waiting',
           'Historical observations', 'Evidence', 'Campaign status', 'Source',
           'Website type', 'Template family', 'Structure group', 'Template match %',
           'Template evidence', 'Detection status', 'Structure checked UTC', 'Structure fingerprint',
           'Template automation', 'Template source', 'Detection notes',
           'Layout status', 'Posting entry point', 'Layout access', 'Layout blueprint',
           'Layout review reason', 'Layout verification')


def now():
    return datetime.now(timezone.utc).isoformat()


@contextmanager
def connect(folder):
    db = sqlite3.connect(Path(folder) / 'attempts.sqlite3', timeout=15)
    try:
        with db:
            db.execute('''CREATE TABLE IF NOT EXISTS attempts (
                id TEXT PRIMARY KEY, domain TEXT NOT NULL, source TEXT NOT NULL,
                outcome TEXT NOT NULL, status TEXT NOT NULL, checked_utc TEXT NOT NULL)''')
            yield db
    finally:
        db.close()


def outcome(record):
    status = record.get('status', 'not_run')
    if status == 'published':
        return 'recovered' if record.get('listing_origin') == 'existing' else 'success'
    if record.get('score_deferred') or status in ('manual_review', 'not_run'):
        return 'deferred'
    if status in ('submitting', 'pending', 'preparing'):
        return 'waiting'
    if status in ('failed', 'rejected', 'unverified', 'uncertain', 'error', 'needs_attention'):
        return 'error'
    return 'deferred'


def seed_history(folder, state):
    """One snapshot per site, not a fabricated count of all historical retries."""
    with connect(folder) as db:
        for host, record in state.items():
            if db.execute('SELECT 1 FROM attempts WHERE domain=?', (host,)).fetchone():
                continue
            # Old needs_attention includes unattempted, missing mappings and user skips.
            result = 'deferred' if record.get('status') == 'needs_attention' else outcome(record)
            db.execute('INSERT OR IGNORE INTO attempts VALUES (?, ?, ?, ?, ?, ?)',
                       ('baseline:' + host, host, 'historical snapshot', result,
                        record.get('status', 'not_run'), record.get('checked_utc', now())))


def begin_attempt(folder, host):
    with connect(folder) as db:
        db.execute('INSERT INTO attempts VALUES (?, ?, ?, ?, ?, ?)',
                   (uuid.uuid4().hex, host, 'attempt', 'waiting', 'preparing', now()))


def ensure_attempt(folder, host):
    with connect(folder) as db:
        exists = db.execute('SELECT 1 FROM attempts WHERE domain=?', (host,)).fetchone()
    if not exists:
        begin_attempt(folder, host)


def sync_history(folder, state):
    """Update the same attempt during receipt/verification/recovery; never add on save."""
    if not (Path(folder) / 'attempts.sqlite3').exists():
        return
    with connect(folder) as db:
        latest = db.execute('SELECT id, domain FROM attempts WHERE rowid IN '
                            '(SELECT MAX(rowid) FROM attempts GROUP BY domain)').fetchall()
        for identity, host in latest:
            if host in state:
                record = state[host]
                result = outcome(record)
                # Preserve uncertainty about unattempted legacy needs_attention records.
                if identity.startswith('baseline:') and record.get('status') == 'needs_attention':
                    result = 'deferred'
                db.execute('UPDATE attempts SET outcome=?, status=?, checked_utc=? WHERE id=?',
                           (result, record.get('status', 'not_run'), record.get('checked_utc', now()), identity))


def ratings(output_root=None, current_folder=None):
    folders = set((output_root or ROOT / 'output').glob('*/attempts.sqlite3'))
    if current_folder and (Path(current_folder) / 'attempts.sqlite3').exists():
        folders.add(Path(current_folder) / 'attempts.sqlite3')
    totals = {}
    for path in sorted(folders):
        with connect(path.parent) as db:
            rows = db.execute('SELECT domain, outcome, source FROM attempts').fetchall()
        for host, result, source in rows:
            score = totals.setdefault(host, {'success': 0, 'error': 0, 'recovered': 0,
                                            'deferred': 0, 'waiting': 0, 'historical': 0})
            score[result] += 1
            score['historical'] += source == 'historical snapshot'
    for score in totals.values():
        n = score['success'] + score['error']
        score['confidence'] = round(100 * (score['success'] + 1) / (n + 2), 1) if n else None
        score['evidence'] = 'Unrated' if not n else 'Low' if n < 5 else 'Medium' if n < 20 else 'High'
    return totals


def ranked_entries(entries, folder):
    scores = ratings(current_folder=folder)
    # Unrated starts at neutral 50; stable ordering breaks ties without claiming certainty.
    return sorted(entries, key=lambda e: -(scores.get(e['domain'], {}).get('confidence')
                                         if scores.get(e['domain'], {}).get('confidence') is not None else 50))


def verified_count(state):
    from mappings.common import publication_domain
    hosts, urls = set(), set()
    for host, record in state.items():
        url = record.get('url', '')
        if record.get('status') != 'published' or not url:
            continue
        try:
            normalized = publication_domain(host, record)
            if domain(url) != normalized or url in urls:
                continue
        except ValueError:
            continue
        hosts.add(normalized)
        urls.add(url)
    return len(hosts)


def configure_target(folder, target):
    if target < 1:
        raise ValueError('Target must be a positive number of distinct websites.')
    path = Path(folder) / 'campaign_settings.json'
    path.write_text(json.dumps({'target': target, 'counting': 'distinct_domains_including_existing'}), encoding='utf-8')


def target_for(folder):
    path = Path(folder) / 'campaign_settings.json'
    return json.loads(path.read_text(encoding='utf-8'))['target'] if path.exists() else DEFAULT_TARGET


def write_progress(folder, state, entries):
    target = target_for(folder)
    count = verified_count(state)
    published = {h for h, record in state.items() if verified_count({h: record})}
    mapped = {e['domain'] for e in entries} & SUPPORTED
    maximum = count + len(mapped - published)
    value = {'target': target, 'verified_websites': count, 'remaining': max(0, target - count),
             'target_met': count >= target, 'pool_websites': len({e['domain'] for e in entries}),
             'mapped_websites': len(mapped), 'maximum_with_current_mappings': maximum,
             'minimum_new_mappings_needed': max(0, target - maximum),
             'counting': 'One verified public backlink per distinct domain; existing matches included'}
    path = Path(folder) / 'progress.json'
    temp = path.with_suffix('.json.tmp')
    temp.write_text(json.dumps(value, indent=2), encoding='utf-8')
    temp.replace(path)


def progress_message(state, target):
    count = verified_count(state)
    return f'{count}/{target} verified websites; {max(0, target - count)} still needed.'


def known_sites():
    sites = {}
    for name in ('sites.csv', 'external_sites.csv', 'live_sites.csv'):
        with (ROOT / 'mappings' / name).open(encoding='utf-8-sig', newline='') as stream:
            for row in csv.DictReader(stream):
                sites.setdefault(row['domain'], {'website': 'https://' + row['domain'], 'enabled': 'Yes',
                                                'notes': row.get('notes', ''), 'source': 'mappings/' + name})
    for host in SUPPORTED:
        sites.setdefault(host, {'website': 'https://' + host, 'enabled': 'Yes', 'notes': '',
                                'source': 'mappings/registry.py'})
    sites['telegra.ph'] = {'website': 'https://telegra.ph/', 'enabled': 'Yes',
                          'notes': 'Official free publishing workflow; live access may vary. No publication test yet.',
                          'source': 'https://telegra.ph/api'}
    return sites


def read_pool(path):
    w = openpyxl.load_workbook(path, data_only=False)
    try:
        if 'Websites' not in w.sheetnames:
            raise ValueError('Website pool needs a Websites sheet.')
        sheet = w['Websites']
        columns = {str(c.value).strip().lower(): c.column - 1 for c in sheet[1] if c.value}
        if not {'website', 'enabled'} <= columns.keys():
            raise ValueError('Website pool needs Website and Enabled columns.')
        rows = []
        seen = set()
        for cells in sheet.iter_rows(min_row=2):
            def value(label):
                if label not in columns:
                    return ''
                cell = cells[columns[label]]
                if cell.data_type == 'f':
                    raise ValueError(f'Websites!{cell.coordinate}: paste a value instead of a formula.')
                return str(cell.value or '').strip()
            url = value('website')
            if not url:
                continue
            host = domain(url)
            enabled = value('enabled').lower()
            if enabled not in ('', 'yes', 'no'):
                raise ValueError(f'Websites row {cells[0].row}: Enabled must be Yes or No.')
            if host in seen:
                raise ValueError('Duplicate website in pool: ' + host)
            seen.add(host)
            rows.append({'website': url, 'domain': host, 'enabled': 'No' if enabled == 'no' else 'Yes',
                         'notes': value('notes'), 'source': value('source')})
        return rows
    finally:
        w.close()


def merge_pool(entries, pool, workbook=None):
    if not Path(pool).exists():
        raise ValueError('Website pool not found. Run python website_pool.py --refresh first.')
    # Per-business Enabled=No remains an exclusion even when the shared pool enables it.
    excluded = set()
    shared_account = {'username': '', 'password': '', 'login_email': ''}
    if workbook:
        w = openpyxl.load_workbook(workbook, data_only=False)
        try:
            if 'Account' in w.sheetnames:
                from workbook_layout import read_account
                shared_account = read_account(w)
            if 'Sites' in w.sheetnames:
                sheet = w['Sites']
                cols = {str(c.value).strip().lower(): c.column - 1 for c in sheet[1] if c.value}
                if {'listing site', 'enabled'} <= cols.keys():
                    for row in sheet.iter_rows(min_row=2):
                        if str(row[cols['enabled']].value or '').strip().lower() == 'no':
                            excluded.add(domain(row[cols['listing site']].value))
        finally:
            w.close()
    existing = {}
    for entry in entries:
        existing.setdefault(entry['domain'], entry)
    return [{**existing.get(row['domain'], {'row': '', 'domain': row['domain'],
                                           **shared_account}),
             'pool_url': row['website']}
            for row in read_pool(pool) if row['enabled'] == 'Yes' and row['domain'] not in excluded]


def workflow(host):
    if host == 'telegra.ph':
        return 'Free API (live access unverified)'
    if host in DIRECTORIES | CLASSIFIEDS:
        return 'Chrome + Playwright'
    return 'Playwright' if host in SUPPORTED else 'Needs mapping'


def refresh_pool(path=DEFAULT_POOL, output_root=None, campaign=None):
    from template_detector import load_catalog, pool_labels
    from map_websites import layout_labels
    path = Path(path)
    known = known_sites()
    rows = read_pool(path) if path.exists() else [dict(v, domain=h) for h, v in sorted(known.items())]
    templates = load_catalog()
    structures = Counter(templates.get(r['domain'], {}).get('fingerprint', '') for r in rows)
    families = {}
    scores = ratings(output_root)
    state = {}
    if campaign is None:
        paths = list((output_root or ROOT / 'output').glob('*/state.json'))
        if paths:
            campaign = max(paths, key=lambda p: p.stat().st_mtime).parent
    if campaign and (Path(campaign) / 'state.json').exists():
        state = json.loads((Path(campaign) / 'state.json').read_text(encoding='utf-8'))
    w = openpyxl.Workbook()
    sheet = w.active
    sheet.title = 'Websites'
    sheet.append(HEADERS)
    for row in rows:
        host = row['domain']
        s = scores.get(host, {})
        values = [row['website'], row['enabled'], row['notes'], host, workflow(host),
                  s.get('confidence') if s.get('confidence') is not None else 'Unrated',
                  s.get('success', 0), s.get('error', 0), s.get('recovered', 0),
                  s.get('deferred', 0) + s.get('waiting', 0), s.get('historical', 0),
                  s.get('evidence', 'Unrated'), state.get(host, {}).get('status', 'not_run'),
                  row.get('source') or known.get(host, {}).get('source', 'User added')]
        labels = pool_labels(host, templates, structures)
        values.extend(labels)
        values.extend(layout_labels(host))
        group = families.setdefault(labels[1], {'sites': 0, 'mapped': 0, 'observed': 0})
        group['sites'] += 1
        group['mapped'] += host in SUPPORTED
        group['observed'] += bool(templates.get(host, {}).get('fingerprint'))
        sheet.append(values)
        for cell in sheet[sheet.max_row]:
            if isinstance(cell.value, str) and cell.value.startswith(('=', '+', '-', '@')):
                cell.data_type = 's'
    sheet.freeze_panes = 'D2'
    sheet.auto_filter.ref = sheet.dimensions
    for cell in sheet[1]:
        cell.font = Font(bold=True, color='FFFFFF')
        cell.fill = PatternFill('solid', fgColor='24476A')
    for index in range(1, len(HEADERS) + 1):
        sheet.column_dimensions[get_column_letter(index)].width = 55 if index in (1, 3, 14, 24, 25) else 25
    summary = w.create_sheet('Template Groups')
    summary.append(['Template family', 'Pool websites', 'Mapped websites', 'Websites with observed forms', 'Needs mapping'])
    for family, totals in sorted(families.items()):
        summary.append([family, totals['sites'], totals['mapped'], totals['observed'], totals['sites'] - totals['mapped']])
    summary.freeze_panes = 'A2'
    summary.auto_filter.ref = summary.dimensions
    for column in 'ABCDE':
        summary.column_dimensions[column].width = 30
    enabled = DataValidation(type='list', formula1='"Yes,No"')
    sheet.add_data_validation(enabled)
    enabled.add('B2:B1048576')
    instructions = w.create_sheet('Instructions')
    for line in (
        'Campaign status shown for: ' + (Path(campaign).name if campaign else 'No saved campaign'),
        'Add any number of rows on Websites. Edit Website, Enabled and Notes; leave calculated columns alone.',
        'Use one HTTP(S) URL or hostname per row. Blank Enabled means Yes. Duplicate domains are rejected.',
        'Keep this file closed while refreshing. Refresh preserves your Website, Enabled and Notes values.',
        'Business details and credentials stay in the business XLSX, never in this shared list.',
        'Target: 100 distinct websites, including previously verified matching links; multiple posts on one domain count once.',
        'Confidence = 100 * (successes + 1) / (successes + errors + 2). No scored results means Unrated (neutral 50 for ordering).',
        'Evidence: Low <5 scored results, Medium 5-19, High >=20. A rating is an observed workflow score, not an SEO metric.',
        'Only verified public results are successes. Recovered existing links count toward the target, separately from publishing success.',
        'Failed/rejected/unverified/uncertain attempts are errors. Pending, human verification waits, missing credentials and user skips do not lower the score.',
        'Historical observations are one saved result per campaign/site, not all earlier retries. Old needs_attention results are unscored.',
        'Each new attempt is recorded once in output/<campaign>/attempts.sqlite3. Receipt checks and repeated saves update that attempt.',
        'progress.json reports the current mapping capacity and the minimum additional mappings needed; mapped routes may still fail live.',
        'The score combines campaign histories across businesses in this project. More reliable workflows run first.',
        'Needs mapping means no reusable adapter yet. A large pool alone does not guarantee 100 available free publications.',
        'Template family is separate from publishing confidence. A 100% structural match is not a success probability or approval to submit.',
        'Template evidence distinguishes live HTTP, saved forms, configured routes, source references and uninspected sites.',
        'Shared structure means identical control fingerprints within this pool; Single observed structure does not prove a website is unique.',
        'Unknown or unmatched forms need inspection. Login/search forms, JavaScript rendering and CAPTCHA can hide the publishing template.',
        'Template detection never enables a route. Native free-plan, field, category and final-submission guards remain required.',
        'Layout inventory files are in mappings/layouts/<domain>.json. A saved blueprint is not an enabled publishing adapter.',
        'python map_websites.py --browser resumes read-only Chrome inspection and prompts for access checks. It never submits listings.',
        'python template_detector.py refreshes saved form evidence; add --scan for read-only HTTP checks, optionally --only example.com.',
        'Chrome runner handles directory/classified workflows and delegates other mapped platforms to Playwright/API. Human checks may still require attention.',
        'python website_pool.py --refresh',
        'python uivision_publish.py "business.xlsx" --publish',
        'python backlinks.py "business.xlsx" --publish --interactive',
    ):
        instructions.append([line])
    instructions.column_dimensions['A'].width = 155
    path.parent.mkdir(parents=True, exist_ok=True)
    temp = path.with_name(path.stem + '.tmp.xlsx')
    w.save(temp)
    w.close()
    temp.replace(path)
    return len(rows)


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--pool', type=Path, default=DEFAULT_POOL)
    parser.add_argument('--refresh', action='store_true')
    parser.add_argument('--campaign-folder', type=Path)
    args = parser.parse_args()
    for path in (ROOT / 'output').glob('*/state.json'):
        state = json.loads(path.read_text(encoding='utf-8'))
        seed_history(path.parent, state)
        sync_history(path.parent, state)
    count = refresh_pool(args.pool, campaign=args.campaign_folder)
    print(f'{count} websites. Pool: {args.pool}')


if __name__ == '__main__':
    main()
