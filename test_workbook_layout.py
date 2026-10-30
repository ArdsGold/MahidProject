"""Template generation, hour conversion and input validation; no site access."""
import tempfile
import unittest
from pathlib import Path

import openpyxl

from backlinks import load_workbook
from workbook_layout import create_template, template_hours


BUSINESS = {'name': 'Example Gym', 'website': 'https://business.example/',
            'description': 'Approved description', 'postal_code': '00123', 'phone': '+1 555 0100'}


class WorkbookLayoutChecks(unittest.TestCase):
    def test_source_hours_become_separate_ascii_cells_and_round_trip(self):
        hours = 'Monday 3\u20139 PM\nTuesday 15:00-21:00\nFriday Closed\nSunday 12-9 PM'
        with tempfile.TemporaryDirectory() as d:
            path = Path(d) / 'client.xlsx'
            create_template(path, ['rentry.co'], {**BUSINESS, 'hours': hours,
                                               'password': 'NEVER-COPY', 'login_email': 'private@example.com'})
            workbook = openpyxl.load_workbook(path)
            self.assertEqual([workbook['Hours'].cell(2, col).value for col in (2, 3, 4)], ['Open', '15:00', '21:00'])
            self.assertEqual(workbook['Hours']['B6'].value, 'Closed')
            self.assertIsNone(workbook['Hours']['C6'].value)
            self.assertEqual(workbook['Hours']['C8'].value, '12:00')
            self.assertTrue(all(c.value is None for row in workbook['Sites'].iter_rows(min_row=2, min_col=2, max_col=4) for c in row))
            for row in workbook['Hours'].iter_rows(min_row=2, min_col=3, max_col=4):
                for cell in row:
                    if cell.value:
                        self.assertTrue(cell.value.isascii())
            workbook.close()
            business, _ = load_workbook(path)
            self.assertEqual(business['hours'], 'Monday: 15:00-21:00\nTuesday: 15:00-21:00\nFriday: Closed\nSunday: 12:00-21:00')
            self.assertEqual(business['postal_code'], '00123')
            self.assertEqual(business['phone'], BUSINESS['phone'])

    def test_reordered_hours_columns_and_rows_preserve_weekday_order(self):
        with tempfile.TemporaryDirectory() as d:
            path = Path(d) / 'client.xlsx'
            create_template(path, ['rentry.co'], BUSINESS)
            workbook = openpyxl.load_workbook(path)
            sheet = workbook['Hours']
            sheet.delete_rows(1, sheet.max_row)
            sheet.append(['Closes', 'Day', 'Opens', 'Status'])
            sheet.append(['02:00', 'Sunday', '22:00', 'Open'])
            sheet.append([None, 'Monday', None, 'Closed'])
            workbook.save(path)
            workbook.close()
            business, _ = load_workbook(path)
            self.assertEqual(business['hours'], 'Monday: Closed\nSunday: 22:00-02:00 (closes next day)')

    def test_invalid_hours_and_renamed_headers_report_cells_before_publishing(self):
        cases = [
            ({'B2': 'Open', 'C2': '3-9 PM', 'D2': '21:00'}, 'Hours!C2'),
            ({'B2': 'Open', 'C2': '15:00', 'D2': '25:00'}, 'Hours!D2'),
            ({'B2': 'Open', 'C2': '=TIME(15,0,0)', 'D2': '21:00'}, 'use a value, not a formula'),
            ({'B2': 'Closed', 'C2': '15:00'}, 'clear Opens/Closes'),
            ({'C1': 'Opening'}, 'keep Day, Status, Opens and Closes'),
            ({'B1': 'Day'}, 'duplicate column headers'),
        ]
        with tempfile.TemporaryDirectory() as d:
            path = Path(d) / 'client.xlsx'
            for cells, error in cases:
                with self.subTest(cells=cells):
                    create_template(path, ['rentry.co'], BUSINESS)
                    workbook = openpyxl.load_workbook(path)
                    for address, value in cells.items():
                        workbook['Hours'][address] = value
                    workbook.save(path)
                    workbook.close()
                    with self.assertRaisesRegex(ValueError, error):
                        load_workbook(path)

    def test_blank_hours_stay_unknown_and_ambiguous_source_is_not_guessed(self):
        self.assertEqual(template_hours(''), {})
        self.assertEqual(template_hours('Monday 11 PM-2 AM')['Monday'], ('Open', '23:00', '02:00'))
        self.assertEqual(template_hours('Monday By appointment')['Monday'], ('By appointment', '', ''))
        with self.assertRaisesRegex(ValueError, 'specify AM/PM'):
            template_hours('Monday 11-1 PM')
        with tempfile.TemporaryDirectory() as d:
            path = Path(d) / 'blank.xlsx'
            create_template(path, ['rentry.co'], BUSINESS)
            business, _ = load_workbook(path)
            self.assertEqual(business['hours'], '')
            workbook = openpyxl.load_workbook(path)
            validations = list(workbook['Hours'].data_validations.dataValidation)
            time_validation = next(v for v in validations if v.type == 'custom')
            self.assertEqual(str(time_validation.sqref), 'C2:D8')
            self.assertEqual(time_validation.errorStyle, 'stop')
            self.assertTrue(time_validation.showErrorMessage)
            workbook.close()


if __name__ == '__main__':
    unittest.main()
