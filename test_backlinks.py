"""Offline checks: synthetic XLSX plus browser-served public/preview pages."""
import csv
import json
import os
import tempfile
import unittest
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import patch

import openpyxl
from playwright.sync_api import sync_playwright, TimeoutError

from backlinks import load_workbook, domain, save, verify, NeedsAttention, campaign_lock, publish, manual_review, prepare, check_public_url
from workbook_layout import create_template
from browser_helper import export_queue, read_results, import_results, write_csv, RESULT_FIELDS


class Checks(unittest.TestCase):
    def test_latest_and_search_pages_do_not_count_but_categories_can(self):
        for path in ('/index.php?p=d', '/index.php?q=Example', '/search.php?q=Example',
                     '/latest-links/', '/index.php?view=search'):
            with self.subTest(path=path), self.assertRaises(NeedsAttention):
                check_public_url('https://directory.test' + path, 'directory.test')
        for path in ('/index.php?c=468', '/Sports/', '/Example_123.html'):
            check_public_url('https://directory.test' + path, 'directory.test')

    def test_final_report_labels_errors_without_changing_recovery_states(self):
        statuses = ['failed', 'needs_attention', 'rejected', 'unverified', 'uncertain',
                    'published', 'pending', 'not_run', 'submitting', 'manual_review']
        entries = [{'row': i + 1, 'domain': f'site{i}.test'} for i in range(len(statuses))]
        entries.append({**entries[0], 'row': 99})
        state = {e['domain']: {'status': status, 'reason': 'Recorded reason',
                              **({'url': 'https://published.test/listing'} if status == 'published' else {})}
                 for e, status in zip(entries, statuses)}
        before = json.dumps(state)
        with tempfile.TemporaryDirectory() as d:
            save(state, Path(d), entries)
            with (Path(d) / 'results.csv').open(encoding='utf-8-sig', newline='') as f:
                rows = list(csv.DictReader(f))
            self.assertEqual([r['status'] for r in rows[:5]], ['error'] * 5)
            self.assertEqual([r['status'] for r in rows[5:10]], statuses[5:])
            self.assertEqual([r['workflow_status'] for r in rows[:10]], statuses)
            self.assertEqual(rows[-1]['status'], 'error')
            self.assertEqual(rows[-1]['duplicate'], 'True')
            self.assertTrue(all(r['reason'] == 'Recorded reason' for r in rows))
            self.assertEqual((Path(d) / 'published_urls.txt').read_text(), 'https://published.test/listing\n')
            published_host = entries[5]['domain']
            state[published_host]['listing_origin'] = 'existing'
            save(state, Path(d), entries)
            with (Path(d) / 'results.csv').open(encoding='utf-8-sig', newline='') as f:
                existing_row = list(csv.DictReader(f))[5]
            self.assertEqual(existing_row['listing_origin'], 'existing')
            self.assertEqual(existing_row['status'], 'published')
            self.assertEqual(existing_row['published_url'], 'https://published.test/listing')
            del state[published_host]['listing_origin']
        self.assertEqual(json.dumps(state), before)

    def test_privatebin_requires_a_complete_viewer_link_and_rejects_deletion_links(self):
        viewer = 'https://privatebin.net/?0123456789abcdef#' + 'A' * 43
        check_public_url(viewer, 'privatebin.net')
        for invalid in (
            'https://privatebin.net/',
            viewer.split('#')[0],
            viewer.replace('?', '?pasteid='),
            viewer.replace('#', '&deletetoken=secret#'),
            viewer.replace('https:', 'http:'),
            viewer.replace('privatebin.net', 'other.example'),
            viewer.replace('A' * 43, 'not-a-key'),
        ):
            with self.subTest(url=invalid), self.assertRaises(NeedsAttention):
                check_public_url(invalid, 'privatebin.net')

    def test_helper_export_is_public_deduplicated_and_does_not_change_state(self):
        business = {'name': 'Example Business', 'website': 'https://business.example/',
                    'description': 'Approved, with "quotes"\nand a newline', 'phone': '+1 012 345 6789',
                    'hours': 'Monday: 09:00-17:00'}
        entries = [{'row': n, 'domain': host, 'username': 'private-user', 'password': 'secret-password',
                    'login_email': 'private@example.com'}
                   for n, host in enumerate(['pastelink.net', 'pastelink.net', 'rentry.co', 'fixture.test'], 1)]
        state = {'pastelink.net': {'status': 'manual_review'},
                 'rentry.co': {'status': 'published', 'url': 'https://rentry.co/already'}}
        before = json.dumps(state)
        with tempfile.TemporaryDirectory() as d:
            folder = Path(d)
            helper, count = export_queue(business, entries, state, folder)
            self.assertEqual(count, 2)
            self.assertEqual(json.dumps(state), before)
            self.assertFalse((folder / 'state.json').exists())
            with (helper / 'helper_queue.csv').open(encoding='utf-8', newline='') as source:
                rows = list(csv.DictReader(source))
            self.assertEqual(rows[0]['action'], 'review')
            self.assertEqual(rows[0]['description'], business['description'])
            self.assertEqual(rows[0]['phone'], business['phone'])
            self.assertEqual(rows[0]['hours'], business['hours'])
            self.assertEqual(rows[1]['action'], 'new')
            exported = ''.join(path.read_text(encoding='utf-8') for path in helper.iterdir())
            for private in ('private-user', 'secret-password', 'private@example.com'):
                self.assertNotIn(private, exported)
            fill = json.loads((helper / 'Pastelink_Fill.json').read_text(encoding='utf-8'))
            self.assertFalse(any(cmd['Command'] in ('click', 'submit') for cmd in fill['Commands']))
            self.assertTrue(any(cmd['Command'] == 'type' and cmd['Target'] == 'id=textbody' for cmd in fill['Commands']))
            state['pastelink.net']['status'] = 'published'
            export_queue(business, entries, state, folder)
            self.assertFalse((helper / 'Pastelink_Fill.json').exists())

    def test_helper_rejects_wrong_campaign_private_urls_and_conflicting_rows(self):
        entries = [{'row': 1, 'domain': 'pastelink.net'}]
        with tempfile.TemporaryDirectory() as d:
            path = Path(d) / 'results.csv'
            valid = ['campaign-one', 'pastelink.net', 'candidate', 'https://pastelink.net/public']
            write_csv(path, [valid, valid])
            self.assertEqual(read_results(path, 'campaign-one', entries), [tuple(valid[1:])])
            write_csv(path, [RESULT_FIELDS, valid])
            self.assertEqual(read_results(path, 'campaign-one', entries), [tuple(valid[1:])])
            for bad in [
                ['other-campaign', *valid[1:]],
                ['campaign-one', 'unknown.test', *valid[2:]],
                ['campaign-one', 'pastelink.net', 'published', valid[3]],
                ['campaign-one', 'pastelink.net', 'candidate', 'https://other.test/public'],
                ['campaign-one', 'pastelink.net', 'candidate', 'https://pastelink.net/edit/secret'],
                ['campaign-one', 'pastelink.net', 'candidate', 'https://pastelink.net/'],
                ['campaign-one', 'pastelink.net', 'pending', valid[3]],
                ['campaign-one', 'pastelink.net', 'skip', ''],  # Conflicts with candidate above.
            ]:
                write_csv(path, [RESULT_FIELDS, valid, bad])
                with self.assertRaises((ValueError, NeedsAttention)):
                    read_results(path, 'campaign-one', entries)

    def test_helper_import_verifies_retains_failures_and_preserves_completed_work(self):
        business = {'name': 'Example Business', 'website': 'https://business.example/'}
        entries = [{'row': 1, 'domain': 'fixture.test'}, {'row': 2, 'domain': 'pastelink.net'},
                   {'row': 3, 'domain': 'rentry.co'}]
        with tempfile.TemporaryDirectory() as d, sync_playwright() as p:
            folder = Path(d)
            browser = p.chromium.launch()
            class FixtureBrowser:
                def new_context(self):
                    ctx = browser.new_context()
                    ctx.route('**/*', lambda route: route.fulfill(content_type='text/html', body=
                        '<h1>Example Business</h1><a href="https://business.example/">Website</a>'))
                    return ctx
            state = {'rentry.co': {'status': 'published', 'url': 'https://rentry.co/previous'}}
            # Offline browser verifies the public page, then a blocked verifier keeps its candidate.
            import_results([('fixture.test', 'candidate', 'https://fixture.test/public')], business,
                           entries, state, folder, FixtureBrowser())
            self.assertEqual(state['fixture.test']['status'], 'published')
            with patch('browser_helper.verify', side_effect=NeedsAttention('Site access challenge')):
                import_results([('pastelink.net', 'candidate', 'https://pastelink.net/public')],
                               business, entries, state, folder, browser)
            self.assertEqual(state['pastelink.net']['status'], 'unverified')
            self.assertEqual(state['pastelink.net']['candidate_url'], 'https://pastelink.net/public')
            with patch('browser_helper.verify') as check:
                import_results([('rentry.co', 'candidate', 'https://rentry.co/different'),
                                ('pastelink.net', 'pending', '')], business, entries, state, folder, browser)
                check.assert_not_called()
            self.assertEqual(state['rentry.co']['url'], 'https://rentry.co/previous')
            self.assertEqual(state['pastelink.net']['status'], 'pending')
            self.assertEqual(state['pastelink.net']['candidate_url'], 'https://pastelink.net/public')
            self.assertEqual(set((folder / 'published_urls.txt').read_text().splitlines()),
                             {'https://rentry.co/previous', 'https://fixture.test/public'})
            browser.close()

    def test_classifieds_fills_correct_location_and_disables_paid_promotions(self):
        with sync_playwright() as p:
            browser = p.chromium.launch()
            page = browser.new_page()
            requests = []
            def fixture(route):
                requests.append(route.request.method)
                url = route.request.url
                if 'stage=form' in url:
                    body = '''<input id="adtitle"><textarea name="addesc"></textarea><input id="email">
                    <select name="arealist"><option value="">Select</option><option value="stl">st louis</option></select><input name="area">
                    <select name="promote[featured]"><option value="paid">15 days ($3)</option><option value="0">None</option></select>
                    <select name="promote[extended]"><option value="paid">30 days ($5)</option><option value="0">None</option></select>
                    <input type="checkbox" name="agree"><button type="submit">Post Now</button>'''
                elif 'stage=category' in url:
                    body = '<a href="/index.php?stage=form">Health &amp; Fitness</a>'
                elif 'stage=state' in url:
                    body = '<a href="/index.php?stage=category">Services</a>'
                else:
                    body = '<a href="/index.php?stage=state">Missouri</a>'
                route.fulfill(content_type='text/html', body=body)
            page.route('https://foldads.com/**', fixture)
            b = {'name': 'Example Gym', 'description': 'Approved description', 'category': 'Sports',
                 'address': '10 Main St, St. Louis, MO 63146, United States',
                 'website': 'https://business.example/', 'email': 'public@example.com'}
            selector, _ = prepare(page, 'foldads.com', b, {}, {}, False)
            self.assertEqual(page.locator('[name="arealist"]').input_value(), 'stl')
            self.assertEqual(page.locator('[name="promote[featured]"]').input_value(), '0')
            self.assertEqual(page.locator('[name="promote[extended]"]').input_value(), '0')
            self.assertIn(b['website'], page.locator('[name="addesc"]').input_value())
            self.assertEqual(page.locator(selector).count(), 1)
            self.assertTrue(all(method == 'GET' for method in requests))
            browser.close()

    def test_open_excel_report_pauses_without_creating_another_csv(self):
        with tempfile.TemporaryDirectory() as d:
            folder = Path(d)
            original = folder / 'results.csv'
            original.write_text('Open in Excel', encoding='utf-8')
            replace = os.replace
            def locked_replace(source, destination):
                if Path(destination) == original:
                    raise PermissionError('File is open in Excel')
                replace(source, destination)
            state = {'rentry.co': {'status': 'published', 'url': 'https://rentry.co/verified'}}
            with patch('backlinks.os.replace', side_effect=locked_replace), patch('backlinks.operator_input', return_value='stop') as prompt:
                with self.assertRaises(KeyboardInterrupt):
                    save(state, folder, [{'domain': 'rentry.co', 'row': 17}])
                save(state, folder, [{'domain': 'rentry.co', 'row': 17}])  # finally cleanup must not prompt again.
                prompt.assert_called_once()
            self.assertEqual(original.read_text(), 'Open in Excel')
            self.assertEqual(list(folder.glob('*.csv')), [original])
            self.assertEqual(json.loads((folder / 'state.json').read_text())['rentry.co']['status'], 'published')
            save(state, folder, [{'domain': 'rentry.co', 'row': 17}])
            self.assertIn('https://rentry.co/verified', original.read_text(encoding='utf-8-sig'))

    def test_report_retry_updates_original_after_excel_unlocks(self):
        import backlinks
        with tempfile.TemporaryDirectory() as d:
            folder = Path(d)
            original_write = backlinks.atomic_write
            blocked = True
            def write(path, value):
                if Path(path).name == 'results.csv' and blocked:
                    raise PermissionError('Excel lock')
                original_write(path, value)
            def unlock(message):
                nonlocal blocked
                blocked = False
                return ''
            with patch('backlinks.atomic_write', side_effect=write), patch('backlinks.operator_input', side_effect=unlock) as prompt:
                save({'rentry.co': {'status': 'published', 'url': 'https://rentry.co/verified'}}, folder, [{'row': 1, 'domain': 'rentry.co'}])
                prompt.assert_called_once()
            self.assertEqual([p.name for p in folder.glob('*.csv')], ['results.csv'])
            self.assertIn('https://rentry.co/verified', (folder / 'results.csv').read_text())

    def test_manual_review_counts_only_verified_public_results(self):
        with tempfile.TemporaryDirectory() as d:
            folder = Path(d)
            b = {'name': 'Example Business', 'website': 'https://business.example/', 'description': 'Approved'}
            page = SimpleNamespace(url='https://rentry.co/new-public-page')
            entries = [{'domain': 'rentry.co', 'row': 17}]
            state = {}
            verified = {'status': 'published', 'url': page.url, 'reason': 'Public check passed'}
            with patch('builtins.input', side_effect=['https://other.test/page', 'current']), patch('backlinks.verify', return_value=verified) as check:
                manual_review(None, page, 'rentry.co', b, state, folder, entries, 'Complete this page')
            check.assert_called_once_with(None, page.url, b, 'rentry.co')
            self.assertEqual((folder / 'published_urls.txt').read_text().strip(), page.url)
            state = {}
            with patch('builtins.input', return_value='pending'), patch('backlinks.verify') as check:
                manual_review(None, page, 'rentry.co', b, state, folder, entries, 'Review approval')
            check.assert_not_called()
            self.assertEqual(state['rentry.co']['status'], 'pending')
            self.assertEqual((folder / 'published_urls.txt').read_text(), '')

    def test_standard_layout_reads_labels_hours_and_enabled_sites(self):
        with tempfile.TemporaryDirectory() as d:
            path = Path(d) / 'input.xlsx'
            create_template(path, ['rentry.co', 'postheaven.net'], {
                'name': 'Example Business', 'website': 'https://business.example/location/',
                'description': 'Approved description', 'phone': '+1 012 345 6789',
                'street': '10 Main St', 'address_line_2': 'Unit 2', 'city': 'Example City',
                'region': 'NJ', 'postal_code': '00123', 'country': 'United States',
            })
            w = openpyxl.load_workbook(path)
            s = w['Business']
            # Move fields and site columns: identification must not depend on their positions.
            first, last = [c.value for c in s[2]], [c.value for c in s[3]]
            for col, value in enumerate(last, 1):
                s.cell(2, col, value)
            for col, value in enumerate(first, 1):
                s.cell(3, col, value)
            sites = w['Sites']
            sites.delete_rows(1, sites.max_row)
            sites.append(['Enabled', 'Login email', 'Listing site', 'Password', 'Username'])
            sites.append(['Yes', 'login@example.com', 'rentry.co', ' secret-password ', 'example-user'])
            sites.append(['No', '', 'postheaven.net', '', ''])
            w['Hours']['B2'] = 'Open'
            from datetime import time
            w['Hours']['C2'], w['Hours']['D2'] = time(9, 0), '17:30'
            w['Hours']['B8'] = 'Closed'
            w.active = w.sheetnames.index('Instructions')
            w.save(path)
            w.close()
            b, entries = load_workbook(path)
            self.assertEqual(b['phone'], '+1 012 345 6789')
            self.assertEqual(b['postal_code'], '00123')
            self.assertEqual(b['address'], '10 Main St Unit 2, Example City, NJ 00123, United States')
            self.assertEqual(b['hours'], 'Monday: 09:00-17:30\nSunday: Closed')
            self.assertEqual(len(entries), 1)
            self.assertEqual(entries[0]['password'], ' secret-password ')
            self.assertEqual(entries[0]['domain'], 'rentry.co')

    def test_standard_layout_rejects_incomplete_data_and_formulas(self):
        with tempfile.TemporaryDirectory() as d:
            path = Path(d) / 'input.xlsx'
            create_template(path, ['rentry.co'])
            with self.assertRaisesRegex(ValueError, 'required field name'):
                load_workbook(path)
            create_template(path, ['rentry.co'], {'name': 'Example Business', 'website': 'https://business.example/', 'description': 'Approved'})
            w = openpyxl.load_workbook(path)
            w['Hours']['B2'], w['Hours']['C2'] = 'Open', '09:00'
            w.save(path)
            w.close()
            with self.assertRaisesRegex(ValueError, 'Hours!D2'):
                load_workbook(path)
            w = openpyxl.load_workbook(path)
            w['Hours']['B2'], w['Hours']['C2'] = None, None
            w['Business']['B2'] = '=1+1'
            w.save(path)
            w.close()
            with self.assertRaisesRegex(ValueError, 'use a value, not a formula'):
                load_workbook(path)

    def test_resume_never_repeats_completed_or_inflight_submissions(self):
        class NoBrowser:
            def new_context(self):
                raise AssertionError('Resume must not submit again')
        with tempfile.TemporaryDirectory() as d:
            folder = Path(d) / 'campaign'
            folder.mkdir()
            entries = [{'domain': h, 'row': i + 17, 'username': '', 'password': '', 'login_email': ''}
                       for i, h in enumerate(['rentry.co', 'postheaven.net'])]
            state = {'rentry.co': {'status': 'published', 'url': 'https://rentry.co/already-created'},
                     'postheaven.net': {'status': 'submitting'}}
            with patch('backlinks.ROOT', Path(d)):
                with self.assertRaises(NeedsAttention):
                    publish(NoBrowser(), {'name': 'Example Business'}, entries, state, folder,
                            SimpleNamespace(only=None, interactive=False))
            self.assertEqual(state['postheaven.net']['status'], 'submitting')

    def test_workbook_and_resume_exports(self):
        with tempfile.TemporaryDirectory() as d:
            folder = Path(d)
            p = folder / 'business.xlsx'
            w = openpyxl.Workbook()
            s = w.active
            s['B1'] = 'Example Business'
            s.append(['Website', 'https://business.example/location/'])
            s.append(['Description:', 'Approved description only'])
            s.append(['Listing Sites', 'Live Listing Links', 'Status', 'Username', 'Password', 'Mail'])
            s.append(['HTTP://WWW.RENTRY.CO/', '', 'Live', 'private-user', 'secret-password', 'private@example.com'])
            s.append(['rentry.co', '', 'Live'])
            w.save(p)
            w.close()
            b, entries = load_workbook(p)
            self.assertEqual(b['description'], 'Approved description only')
            self.assertEqual([e['domain'] for e in entries], ['rentry.co', 'rentry.co'])
            self.assertEqual(entries[0]['password'], 'secret-password')
            state = {'rentry.co': {'status': 'unverified', 'candidate_url': 'https://rentry.co/not-public'}}
            save(state, folder, entries)
            self.assertEqual((folder / 'published_urls.txt').read_text(), '')
            state['rentry.co'] = {'status': 'published', 'url': 'https://rentry.co/public'}
            save(state, folder, entries)
            self.assertEqual((folder / 'published_urls.txt').read_text(), 'https://rentry.co/public\n')
            exported = (folder / 'results.csv').read_text(encoding='utf-8-sig') + (folder / 'state.json').read_text()
            self.assertNotIn('secret-password', exported)
            self.assertNotIn('private-user', exported)
            with (folder / 'results.csv').open(encoding='utf-8-sig', newline='') as source:
                report_rows = list(csv.DictReader(source))
            self.assertEqual(report_rows[1]['duplicate'], 'True')
            self.assertEqual(report_rows[1]['status'], 'published')
            with campaign_lock(folder / '.lock'):
                with self.assertRaises(ValueError):
                    with campaign_lock(folder / '.lock'):
                        pass
            self.assertFalse((folder / '.lock').exists())
        with self.assertRaises(ValueError):
            domain('https://user:password@rentry.co')

    def test_public_link_verification(self):
        business = {'name': 'Example Business', 'website': 'https://business.example/location/'}
        with sync_playwright() as p:
            browser = p.chromium.launch()
            class FixtureBrowser:
                def __init__(self, body):
                    self.body = body
                def new_context(self):
                    ctx = browser.new_context()
                    ctx.route('https://fixture.test/**', lambda route: route.fulfill(content_type='text/html', body=self.body))
                    return ctx
            public = '<meta name="robots" content="noindex"><h1>Example Business</h1><a href="https://business.example/location/" rel="nofollow">Website</a>'
            result = verify(FixtureBrowser(public), 'https://fixture.test/listing', business, 'fixture.test', 300)
            self.assertEqual(result['status'], 'published')
            self.assertTrue(result['nofollow'])
            self.assertTrue(result['noindex'])
            styled = '<style>h1{text-transform:uppercase}</style>' + public
            self.assertEqual(verify(FixtureBrowser(styled), 'https://fixture.test/listing', business, 'fixture.test', 300)['status'], 'published')
            for body in [
                '<h1>Example Business</h1><p>https://business.example/location/</p>',
                '<h1>Example Business</h1><a href="https://business.example/other/">Other page</a>',
                '<h1>Example Business</h1><a style="display:none" href="https://business.example/location/">Hidden</a>',
                '<h1>Sign in to see this listing</h1>',
                '<h1>Another Business</h1><a href="https://business.example/location/">Website</a>',
            ]:
                with self.assertRaises(TimeoutError):
                    verify(FixtureBrowser(body), 'https://fixture.test/listing', business, 'fixture.test', 300)
            with self.assertRaises(NeedsAttention):
                verify(FixtureBrowser(public), 'https://fixture.test/edit/private', business, 'fixture.test', 300)
            with self.assertRaises(NeedsAttention):
                verify(FixtureBrowser(public), 'https://other.test/listing', business, 'fixture.test', 300)
            browser.close()


if __name__ == '__main__':
    unittest.main()
