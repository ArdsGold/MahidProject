"""Paste-ready client workbook integration checks; no network or publication."""
import tempfile
import unittest
from pathlib import Path

import openpyxl
from playwright.sync_api import sync_playwright

from backlinks import load_workbook
from workbook_layout import create_simple_template, create_template
from website_pool import merge_pool
from site_access import shared_credentials
from mappings.categories import category_choice
from mappings.common import content_parts
from uivision_publish import make_job

BUSINESS = {'name':'Example Battery','address':'1929 S State St Unit 3, Chicago, IL 60616',
            'phone':'+1 555 0100','website':'https://battery.example/chicago/',
            'description':'Approved mobile battery service.','category':'Car battery store',
            'hours':'Monday 7\u202fAM\u20138\u202fPM\nTuesday Closed\nSunday 10 PM-2 AM',
            'keywords':'battery, roadside','category_fallbacks':['Roadside assistance','Automotive services']}
ACCOUNT = {'username':'private-user','login_email':'private@example.test','password':' private-pass '}

class SimpleWorkbookChecks(unittest.TestCase):
    def test_first_eight_rows_match_original_and_no_repeated_credentials(self):
        with tempfile.TemporaryDirectory() as d:
            p=Path(d)/'client.xlsx';create_simple_template(p,['one.example','two.example'],BUSINESS,ACCOUNT)
            w=openpyxl.load_workbook(p)
            self.assertEqual(w.sheetnames,['Business','Account','Sites','Overrides'])
            self.assertEqual([w['Business'].cell(i,1).value for i in range(1,9)],
                             ['Business name','Address','Phone','Website','Description','Category','Business hours','Keywords'])
            self.assertEqual(w['Sites'].max_column,2)
            for value in ACCOUNT.values():
                self.assertEqual(sum(c.value==value for sh in w for row in sh for c in row),1)
            w.close()
            b,e=load_workbook(p)
            self.assertEqual(b['address'],BUSINESS['address'])
            self.assertEqual((b['street'],b['city'],b['region'],b['postal_code']),('1929 S State St Unit 3','Chicago','IL','60616'))
            self.assertEqual(b['hours'],'Monday: 07:00-20:00\nTuesday: Closed\nSunday: 22:00-02:00 (closes next day)')
            self.assertEqual(b['category_fallbacks'],BUSINESS['category_fallbacks'])
            self.assertEqual(e[0]['password'],ACCOUNT['password'])
            public=str(content_parts(b))
            self.assertTrue(all(value not in public for value in ACCOUNT.values()))
            self.assertEqual(make_job(b,'one.example')['category_fallbacks'],BUSINESS['category_fallbacks'])

    def test_paste_original_labels_with_colons_and_keep_text_values(self):
        with tempfile.TemporaryDirectory() as d:
            p=Path(d)/'client.xlsx';create_simple_template(p,['one.example'])
            w=openpyxl.load_workbook(p)
            values=[BUSINESS[k] for k in ['name','address','phone','website','description','category','hours','keywords']]
            for i,v in enumerate(values,1):
                w['Business'].cell(i,1).value=str(w['Business'].cell(i,1).value)+':'
                w['Business'].cell(i,2).value=v
            w.save(p);w.close()
            b,_=load_workbook(p);self.assertEqual(b['phone'],BUSINESS['phone']);self.assertEqual(b['name'],BUSINESS['name'])

    def test_shared_account_applies_to_added_pool_sites_and_overrides_only_one_site(self):
        with tempfile.TemporaryDirectory() as d:
            p=Path(d)/'client.xlsx';pool=Path(d)/'pool.xlsx'
            create_simple_template(p,['one.example'],BUSINESS,ACCOUNT)
            w=openpyxl.load_workbook(p);w['Overrides'].append(['two.example','different-user','','different-pass']);w.save(p);w.close()
            w=openpyxl.Workbook();s=w.active;s.title='Websites';s.append(['Website','Enabled']);s.append(['one.example','Yes']);s.append(['two.example','Yes']);s.append(['three.example','Yes']);w.save(pool);w.close()
            _,e=load_workbook(p);merged={r['domain']:r for r in merge_pool(e,pool,p)}
            self.assertEqual(merged['three.example']['password'],ACCOUNT['password'])
            self.assertEqual(merged['two.example']['password'],'different-pass')
            self.assertEqual(merged['two.example']['login_email'],ACCOUNT['login_email'])
            self.assertEqual(shared_credentials(p)['username'],ACCOUNT['username'])
            self.assertEqual(shared_credentials(p,'two.example')['username'],'different-user')
            self.assertEqual(shared_credentials(p,'three.example')['password'],ACCOUNT['password'])

    def test_disabled_site_cannot_be_reenabled_by_override(self):
        with tempfile.TemporaryDirectory() as d:
            p=Path(d)/'client.xlsx';pool=Path(d)/'pool.xlsx';create_simple_template(p,['one.example'],BUSINESS,ACCOUNT)
            w=openpyxl.load_workbook(p);w['Sites']['B2']='No';w['Overrides'].append(['one.example','other','','']);w.save(p);w.close()
            w=openpyxl.Workbook();s=w.active;s.title='Websites';s.append(['Website','Enabled']);s.append(['one.example','Yes']);w.save(pool);w.close()
            _,e=load_workbook(p);self.assertEqual(merge_pool(e,pool,p),[])

    def test_invalid_hours_formulas_duplicate_sites_and_overrides_fail_early(self):
        cases=[('Business','B7','Monday 11-1 PM','specify AM/PM'),('Business','B7','Monday 99:00-20:00','Business hours'),
               ('Business','B7','Monday 7 AM-8 PM\nMonday Closed','duplicate'),
               ('Business','B5','=1+1','use a value'),('Account','B4','=1+1','use a value'),
               ('Sites','A3','https://www.one.example/path','duplicate website'),
               ('Overrides','B2','orphan-user','enter a Listing site')]
        with tempfile.TemporaryDirectory() as d:
            p=Path(d)/'client.xlsx'
            for sheet,cell,value,error in cases:
                with self.subTest(sheet=sheet,cell=cell,value=value):
                    create_simple_template(p,['one.example'],BUSINESS,ACCOUNT)
                    w=openpyxl.load_workbook(p);w[sheet][cell]=value;w.save(p);w.close()
                    with self.assertRaisesRegex(ValueError,error):load_workbook(p)

    def test_empty_sites_uses_pool_and_account_never_copied_implicitly(self):
        with tempfile.TemporaryDirectory() as d:
            p=Path(d)/'client.xlsx';create_simple_template(p,business={**BUSINESS,**ACCOUNT})
            _,entries=load_workbook(p);self.assertEqual(entries,[])
            w=openpyxl.load_workbook(p)
            self.assertTrue(all(w['Account'].cell(i,2).value is None for i in range(2,5)))
            w.close()

    def test_legacy_shared_credentials_still_apply_to_a_new_pool_host(self):
        with tempfile.TemporaryDirectory() as d:
            p=Path(d)/'legacy.xlsx';create_template(p,['one.example'],BUSINESS)
            w=openpyxl.load_workbook(p)
            w['Sites']['B2']=ACCOUNT['username'];w['Sites']['C2']=ACCOUNT['password'];w['Sites']['D2']=ACCOUNT['login_email']
            w.save(p);w.close()
            self.assertEqual(shared_credentials(p,'new.example')['username'],ACCOUNT['username'])

    def test_fallback_categories_are_ordered_and_only_actual_unambiguous_options_match(self):
        with sync_playwright() as p:
            browser=p.chromium.launch();page=browser.new_page()
            options=[{'label':'Automotive services','value':'auto'},{'label':'Roadside assistance','value':'road'}]
            self.assertEqual(category_choice(page,options,BUSINESS)['value'],'road')
            self.assertEqual(category_choice(page,[options[0]],BUSINESS)['value'],'auto')
            self.assertEqual(category_choice(page,options+[{'label':'Car battery store','value':'battery'}],BUSINESS)['value'],'battery')
            self.assertIsNone(category_choice(page,[options[1],{**options[1],'value':'road2'},options[0]],BUSINESS))
            self.assertIsNone(category_choice(page,[{'label':'Veterinary services','value':'vet'}],BUSINESS))
            browser.close()

if __name__=='__main__':unittest.main()
