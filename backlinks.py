"""Local XLSX -> Playwright publisher. See README.md before running."""
import argparse
import csv
import concurrent.futures
import hashlib
import io
import json
import os
import re
import secrets
import sys
import time
import webbrowser
from contextlib import contextmanager
from datetime import datetime, timezone
from pathlib import Path
from urllib.parse import urlsplit, urljoin, parse_qs

import openpyxl
from playwright.sync_api import sync_playwright, TimeoutError as BrowserTimeout

from inspect_sites import inspect
from workbook_layout import read_standard
from website_pool import (DEFAULT_POOL, DEFAULT_TARGET, merge_pool, ranked_entries,
                          begin_attempt, ensure_attempt, seed_history, sync_history, verified_count,
                          configure_target, target_for, write_progress, progress_message,
                          refresh_pool)

from mappings.registry import CLASSIFIEDS, DIRECTORIES, US_STATES, SUPPORTED
from mappings.common import (NeedsAttention, InspectionRequired, RefillRequested, domain, content, go,
                             challenge, wait_for_challenge, fill_text)
from mappings.playwright import prepare
from mappings.discovery import directory_candidates
from template_detector import FORM_STRUCTURE_JS, observe_browser
from operator_notifications import operator_input, check_cancelled

ROOT = Path(__file__).resolve().parent
PAUSED_REPORTS = set()


def text(cell):
    # Cached formula values and cell text are data; workbook instructions are never executed.
    return str(cell.value or '').strip()


def load_workbook(path, sheet=None):
    w = openpyxl.load_workbook(path, data_only=False)
    try:
        if {'Business', 'Sites'} <= set(w.sheetnames) and sheet in (None, 'Business'):
            return read_standard(w, domain)
        if 'Business' in w.sheetnames and text(w['Business']['A1']).lower() == 'field':
            raise ValueError('Standard template requires Business and Sites sheets; omit --sheet or use --sheet Business.')
        w.close()
        w = openpyxl.load_workbook(path, data_only=True)
        s = w[sheet] if sheet else w.active
        business = {'name': text(s['B1'])}
        aliases = {'address': 'address', 'phone': 'phone', 'website': 'website',
                   'description': 'description', 'category': 'category', 'catagory': 'category',
                   'business hours': 'hours', 'hours': 'hours', 'keywords': 'keywords',
                   **{key: key for key in ('facebook', 'instagram', 'linkedin', 'youtube', 'tiktok', 'twitter', 'timezone', 'services', 'short_description')}, 'business email': 'email', 'logo': 'logo'}
        header = None
        for row in s.iter_rows():
            label = text(row[0]).lower().rstrip(':').strip()
            if label == 'listing sites':
                header = row[0].row
                break
            if label in aliases:
                c = row[1]
                business[aliases[label]] = c.hyperlink.target if c.hyperlink and label in ('website', 'logo', 'facebook', 'instagram', 'linkedin', 'youtube', 'tiktok', 'twitter') else text(c)
        for key in ('name', 'website', 'description'):
            if not business.get(key):
                raise ValueError('Missing business field: ' + key)
        if urlsplit(business['website']).scheme not in ('http', 'https'):
            raise ValueError('Business website must include http:// or https://.')
        if not header:
            raise ValueError('Missing Listing Sites header.')
        business['description'] = business['description'].strip('"')
        # User confirmed legacy hour-range separators should be standard hyphens.
        # Other damaged characters remain unresolved and are reported by the formatter.
        if business.get('hours'):
            business['hours'] = re.sub(r'(?<=\d)[\ufffd\u2013\u2014](?=\d)', '-', business['hours'])
        entries = []
        for row in s.iter_rows(min_row=header + 1):
            if text(row[0]):
                entries.append({'row': row[0].row, 'domain': domain(text(row[0])),
                                'username': text(row[3]) if len(row) > 3 else '',
                                'password': text(row[4]) if len(row) > 4 else '',
                                'login_email': text(row[5]) if len(row) > 5 else ''})
        if not entries:
            raise ValueError('No listing sites found.')
        return business, entries
    finally:
        w.close()


def stamp():
    return datetime.now(timezone.utc).isoformat()


def campaign_folder(business, campaign='default'):
    key = hashlib.sha256((business['website'].rstrip('/') + '\n' + business['name'] + '\n' + campaign).encode()).hexdigest()[:12]
    slug = re.sub('[^a-z0-9]+', '-', business['name'].lower()).strip('-')[:45]
    return ROOT / 'output' / (slug + '-' + key)


def atomic_write(path, value):
    temp = path.with_suffix(path.suffix + '.tmp')
    temp.write_text(value, encoding='utf-8', newline='')
    os.replace(temp, path)


@contextmanager
def campaign_lock(path):
    try:
        fd = os.open(path, os.O_CREAT | os.O_EXCL | os.O_WRONLY)
    except FileExistsError:
        raise ValueError('Campaign is locked. If no publisher is running, delete ' + str(path)) from None
    try:
        os.write(fd, str(os.getpid()).encode())
        os.close(fd)
        yield
    finally:
        path.unlink(missing_ok=True)


def save(state, folder, entries):
    atomic_write(folder / 'state.json', json.dumps(state, indent=2, ensure_ascii=False))
    sync_history(folder, state)
    write_progress(folder, state, entries)
    urls = list(dict.fromkeys(r['url'] for r in state.values() if r.get('status') == 'published'))
    atomic_write(folder / 'published_urls.txt', ''.join(u + '\n' for u in urls))
    f = io.StringIO(newline='')
    writer = csv.writer(f)
    writer.writerow(['spreadsheet_row', 'domain', 'status', 'published_url', 'reason', 'checked_utc', 'nofollow', 'noindex', 'redirected', 'workflow_status', 'duplicate', 'listing_origin', 'omitted_fields', 'content_notes'])
    seen = set()
    listed = {e['domain'] for e in entries}
    report_entries = list(entries) + [{'row': '', 'domain': host} for host in state if host not in listed]
    for e in report_entries:
        r = state.get(e['domain'], {})
        duplicate = e['domain'] in seen
        seen.add(e['domain'])
        workflow_status = r.get('status', 'not_run')
        status = 'error' if workflow_status in ('failed', 'needs_attention', 'needs_mapping', 'rejected', 'unverified', 'uncertain', 'error') else workflow_status
        values = [e['row'], e['domain'], status, r.get('url', '') if workflow_status == 'published' else '', r.get('reason', ''), r.get('checked_utc', ''), r.get('nofollow', ''), r.get('noindex', ''), r.get('redirected', ''), workflow_status, duplicate, r.get('listing_origin', ''), '; '.join(r.get('omitted_fields', [])), r.get('content_notes', '')]
        # Prevent spreadsheet formula interpretation in externally controlled strings.
        writer.writerow(["'" + v if isinstance(v, str) and v.startswith(('=', '+', '-', '@')) else v for v in values])
    report = '\ufeff' + f.getvalue()
    report_path = (folder / 'results.csv').resolve()
    while True:
        try:
            atomic_write(report_path, report)
            PAUSED_REPORTS.discard(report_path)
            break
        except PermissionError:
            # A finally-block save must still persist the journal after the operator pauses.
            if report_path in PAUSED_REPORTS:
                return
            try:
                answer = operator_input('results.csv is open or locked. Close it in Excel, then press Enter to update the same file. Type stop to pause: ').strip().lower()
                if answer in ('stop', 'skip'):
                    raise KeyboardInterrupt
            except (EOFError, KeyboardInterrupt):
                PAUSED_REPORTS.add(report_path)
                print('Paused: close results.csv in Excel before resuming. Progress is saved; no extra CSV was created.', flush=True)
                raise KeyboardInterrupt from None


def check_public_url(url, host):
    if domain(url) != host:
        raise NeedsAttention('Public URL points outside the listing domain')
    parsed = urlsplit(url)
    # PrivateBin viewer links use a paste ID in the query and a public decryption key.
    if (host == 'privatebin.net' and parsed.scheme == 'https' and parsed.path in ('', '/')
            and re.fullmatch(r'[0-9a-f]{16}', parsed.query)
            and re.fullmatch(r'[1-9A-HJ-NP-Za-km-z]{32,44}', parsed.fragment)):
        return
    view = parse_qs(parsed.query).get('view', [''])[0].lower()
    query = {k.lower(): v for k, v in parse_qs(parsed.query, keep_blank_values=True).items()}
    if (any(k in query for k in ('q', 'search', 'query'))
            or query.get('p', [''])[0].lower() in ('d', 'latest', 'new')
            or view in ('search', 'latest')
            or re.search(r'(^|/)(search|latest|latest-links|new-links)([./]|$)', parsed.path, re.I)):
        raise NeedsAttention('Latest Links and Search pages cannot count; use a detail or category page')
    if parsed.path.rstrip('/') in ('', '/submit', '/submit.php', '/login', '/signup') or view in ('post', 'edit', 'selectcity', 'login') or re.search(r'(^|/)edit($|/)|[?&](edit|token|password)=', parsed.path + '?' + parsed.query, re.I):
        raise NeedsAttention('Submission, home, or edit page cannot count as a public listing')


def verify(browser, url, business, host, timeout_ms=15000):
    check_public_url(url, host)
    ctx = browser.new_context()  # Independent cookies: a logged-in preview is not a public listing.
    try:
        page = ctx.new_page()
        response = go(page, url)
        check_public_url(page.url, host)  # Redirects must also satisfy the public-page policy.
        target = business['website']
        redirect_href = ''
        if host == 'ikeanded.com':
            link = page.locator('#PageContent_Website')
            if link.count() == 1:
                href = link.get_attribute('href') or ''
                absolute = urljoin(page.url, href)
                parsed_link = urlsplit(absolute)
                if domain(absolute) == host and parsed_link.path == '/util/urlclick.aspx' and parse_qs(parsed_link.query).get('url') == [target]:
                    check = ctx.request.get(absolute, max_redirects=0, timeout=15000)
                    location = urljoin(absolute, check.headers.get('location', ''))
                    if check.status in (301, 302, 303, 307, 308) and location.rstrip('/') == target.rstrip('/'):
                        redirect_href = absolute
        # Wait for client-rendered public links; exact destination prevents unrelated links counting.
        page.wait_for_function('''({target,name,redirect_href}) => {
          const normalize = u => {const x=new URL(u,location.href);return x.origin.toLowerCase()+x.pathname.replace(/\\/$/,'')+x.search;};
          const normalizedText = value => value.normalize('NFKC').toLowerCase().replace(/\\s+/g, ' ').trim();
          return normalizedText(document.body.innerText).includes(normalizedText(name)) && [...document.querySelectorAll('a[href]')].some(a => {
            try {return (normalize(a.href)===normalize(target) || a.href===redirect_href) && a.getClientRects().length>0;} catch{return false;}
          });
        }''', arg={'target': target, 'name': business['name'], 'redirect_href': redirect_href}, timeout=timeout_ms)
        flags = page.evaluate('''({target,redirect_href}) => {
          const norm = u => {const x=new URL(u,location.href);return x.origin.toLowerCase()+x.pathname.replace(/\\/$/,'')+x.search;};
          const links=[...document.querySelectorAll('a[href]')].filter(a=>{try{return (norm(a.href)===norm(target)||a.href===redirect_href)&&a.getClientRects().length>0;}catch{return false;}});
          return {nofollow:links.every(a=>a.rel.split(/\\s+/).includes('nofollow')), noindex:[...document.querySelectorAll('meta[name="robots"],meta[name="googlebot"]')].some(m=>/noindex/i.test(m.content))};
        }''', {'target': target, 'redirect_href': redirect_href})
        if response and re.search(r'\bnoindex\b', response.headers.get('x-robots-tag', ''), re.I):
            flags['noindex'] = True
        return {'status': 'published', 'url': page.url, 'reason': 'Public page contains business name and clickable website link', 'checked_utc': stamp(), 'redirected': bool(redirect_href), **flags}
    finally:
        ctx.close()


def audit(entries, folder):
    hosts = list(dict.fromkeys(e['domain'] for e in entries))
    results = []
    urls = ['https://' + host + ('/submit.php' if host in DIRECTORIES else '/submit' if host == 'pastelink.net' else '/') for host in hosts]
    # Read-only HTTP discovery is parallel; publishing stays sequential and uses Playwright.
    with concurrent.futures.ThreadPoolExecutor(max_workers=8) as pool:
        for host, item in zip(hosts, pool.map(inspect, urls)):
            item.update(domain=host, checked_utc=stamp(), workflow=host in SUPPORTED)
            results.append(item)
            atomic_write(folder / 'audit.json', json.dumps(results, indent=2, ensure_ascii=False))
            print(host, 'CAPTCHA markup' if item.get('captcha') else 'HTTP ' + str(item.get('status')) if item.get('status') else 'network or HTTP error', flush=True)


def recover_directory(browser, business, host, state, folder, entries, origin):
    """A receipt/duplicate may already have a public detail page; never repost to find it."""
    try:
        candidates = directory_candidates(browser, business, host)
    except Exception:
        return False  # Search access differs from the regular Chrome session.
    for url in candidates:
        try:
            check_public_url(url, host)
        except (ValueError, NeedsAttention):
            continue
        state[host]['candidate_url'] = url
        save(state, folder, entries)  # Preserve discovered URLs before independent verification.
        try:
            result = verify(browser, url, business, host)
        except Exception:
            continue
        result['reason'] += '; public detail page recovered by exact-destination directory search'
        state[host] = {**state[host], **result, 'listing_origin': origin}
        save(state, folder, entries)
        return True
    return False


def manual_review(browser, page, host, b, state, folder, entries, reason):
    """Keep the browser open so the user can finish a missing or interrupted flow."""
    ensure_attempt(folder, host)
    print('\n' + host + ': ' + reason, flush=True)
    print('Complete this site in the browser. If a submission may already exist, find its public page before creating another.', flush=True)
    print('Business: ' + b['name'] + '\nWebsite: ' + b['website'] + '\nDescription: ' + b['description'], flush=True)
    for key in ('address', 'phone', 'email', 'category'):
        if b.get(key):
            print(key.capitalize() + ': ' + b[key], flush=True)
    state.setdefault(host, {}).update(status='manual_review', reason=reason, checked_utc=stamp())
    save(state, folder, entries)
    if page.url == 'about:blank':
        try:
            page.goto(state[host].get('candidate_url', 'https://' + host + '/'), wait_until='domcontentloaded', timeout=25000)
        except Exception:
            print('Page did not load normally. You can retry or navigate in the open browser.', flush=True)
    while True:
        answer = operator_input('Paste the PUBLIC result URL, or type current / browser / pending / skip / stop: ').strip()
        if answer.lower() == 'browser':
            path = '/submit' if host == 'pastelink.net' else '/submit.php' if host in DIRECTORIES else '/'
            webbrowser.open('https://' + host + path)
            print('Opened your regular browser. Complete the site there, then paste its public result URL here.', flush=True)
            continue
        if answer.lower() == 'stop':
            raise KeyboardInterrupt
        if answer.lower() in ('skip', 'pending'):
            state[host].update(status='pending' if answer.lower() == 'pending' else 'needs_attention',
                               score_deferred=answer.lower() == 'skip',
                               reason='User reported awaiting site approval' if answer.lower() == 'pending' else 'User skipped during interactive review', checked_utc=stamp())
            save(state, folder, entries)
            return
        url = page.url if answer.lower() == 'current' else answer
        try:
            check_public_url(url, host)
        except (ValueError, NeedsAttention):
            print('Enter a public listing URL on ' + host + ', excluding login, edit, or home pages.', flush=True)
            continue
        state[host].update(candidate_url=url, status='unverified', reason='Checking user-supplied public URL', checked_utc=stamp())
        save(state, folder, entries)
        try:
            result = verify(browser, url, b, host)
        except Exception:
            state[host]['reason'] = 'Public verification failed during interactive review'
            save(state, folder, entries)
            print('Not verified: the public page must contain the business name and a visible clickable website link. Check the URL or use pending/skip.', flush=True)
            continue
        result['reason'] += '; completed interactively'
        if host == 'ikeanded.com':
            result['reason'] += '; free plan lasts 5 days'
        state[host] = result
        save(state, folder, entries)
        print('VERIFIED ' + result['url'], flush=True)
        return


def review_in_new_page(browser, host, b, state, folder, entries, reason):
    ctx = browser.new_context()
    try:
        manual_review(browser, ctx.new_page(), host, b, state, folder, entries, reason)
    finally:
        ctx.close()


def publish_telegraph(browser, b, entries, state, folder, recovery, recovery_file):
    from mappings.telegraph import call, payload
    host = 'telegra.ph'
    begin_attempt(folder, host)
    started = False
    state[host] = {'status': 'needs_attention', 'score_deferred': True,
                   'reason': 'Preparing Telegraph API request', 'checked_utc': stamp()}
    save(state, folder, entries)
    try:
        values, omitted = payload(b)
        state[host].update(omitted_fields=omitted, content_notes='All available public fields; 64 KB API content limit')
        account = recovery.setdefault(host, {})
        if not account.get('access_token'):
            check_cancelled()
            account['access_token'] = call('createAccount', short_name=b['name'][:32])['access_token']
            atomic_write(recovery_file, json.dumps(recovery, indent=2))
        state[host].update(status='submitting', score_deferred=False,
                           reason='Telegraph createPage started; do not automatically repeat')
        save(state, folder, entries)
        started = True
        check_cancelled()
        result = call('createPage', access_token=account['access_token'], **values)
        check_public_url(result['url'], host)
        state[host].update(status='unverified', candidate_url=result['url'], reason='Telegraph URL awaiting independent verification')
        save(state, folder, entries)
        state[host].update(verify(browser, result['url'], b, host))
    except NeedsAttention as exc:
        state[host].update(status='uncertain' if started else 'needs_attention', score_deferred=False,
                           reason=str(exc), checked_utc=stamp())
    except Exception:
        state[host].update(status='uncertain' if started else 'failed', score_deferred=False,
                           reason='Telegraph API/network or public verification failed; inspect saved result before retrying', checked_utc=stamp())
    finally:
        save(state, folder, entries)


def publish(browser, b, entries, state, folder, args):
    review_interactive = getattr(args, 'manual_review', args.interactive)
    private = ROOT / '.local' / folder.name
    private.mkdir(parents=True, exist_ok=True)
    recovery_file = private / 'recovery.json'
    recovery = json.loads(recovery_file.read_text(encoding='utf-8')) if recovery_file.exists() else {}
    unique = {}
    for e in entries:
        unique.setdefault(e['domain'], e)
    ordered = ranked_entries(list(unique.values()), folder)
    if getattr(args, 'start_with', None):
        ordered.sort(key=lambda e: e['domain'] != args.start_with)
    for entry in ordered:
        check_cancelled()
        if not getattr(args, 'all_sites', False) and verified_count(state) >= target_for(folder):
            break
        host = entry['domain']
        if args.only and host not in args.only:
            continue
        previous = state.get(host, {})
        if previous.get('review_disposition') == 'operator_deferred' or previous.get('status') == 'skipped':
            continue
        if previous.get('status') == 'published':
            continue
        if previous.get('candidate_url'):
            try:
                from mappings.common import publication_domain
                state[host] = {**previous, **verify(browser, previous['candidate_url'], b, publication_domain(host, previous))}
            except Exception:
                previous['status'] = 'unverified'
                previous['reason'] = 'Saved candidate is not a verified public backlink; submission will not be repeated'
                if review_interactive:
                    review_in_new_page(browser, host, b, state, folder, entries, previous['reason'])
            save(state, folder, entries)
            if state[host].get('status') in ('uncertain', 'submitting', 'unverified'):
                raise InspectionRequired(host, 'Public candidate remains unresolved after inspection; queue paused without reposting')
            continue
        if previous.get('status') in ('submitting', 'uncertain', 'pending', 'manual_review', 'rejected'):
            if host in DIRECTORIES and previous.get('status') in ('pending', 'rejected', 'uncertain'):
                if recover_directory(browser, b, host, state, folder, entries,
                                     'existing' if previous.get('status') == 'rejected' else 'new'):
                    continue
            if review_interactive:
                review_in_new_page(browser, host, b, state, folder, entries, 'An earlier submission needs review; no automatic repost will be attempted.')
                continue
            print(host, 'submission already attempted; check before retrying', flush=True)
            if previous.get('status') in ('uncertain', 'submitting'):
                raise InspectionRequired(host, 'Earlier submission needs inspection; queue paused without reposting')
            continue
        if host not in SUPPORTED:
            if review_interactive:
                review_in_new_page(browser, host, b, state, folder, entries, 'No automatic workflow yet. Complete the site in the browser, then provide its public URL.')
                continue
            state[host] = {'status': 'needs_attention', 'reason': 'No implemented submission workflow for this platform; see audit.json', 'checked_utc': stamp()}
            save(state, folder, entries)
            print(host, 'needs_attention: site-specific workflow required', flush=True)
            continue
        if host == 'telegra.ph':
            publish_telegraph(browser, b, entries, state, folder, recovery, recovery_file)
            print(host, state[host]['status'], flush=True)
            continue
        ctx = browser.new_context()
        page = ctx.new_page()
        page.set_default_timeout(12000)
        started = False
        begin_attempt(folder, host)
        state[host] = {'status': 'needs_attention', 'score_deferred': True,
                       'reason': 'Form preparation started; no submission yet', 'checked_utc': stamp()}
        save(state, folder, entries)
        try:
            if host not in recovery:
                slug = re.sub('[^a-z0-9]+', '-', b['name'].lower()).strip('-')[:40] + '-' + secrets.token_hex(4)
                recovery[host] = {'slug': slug, 'edit_code': secrets.token_urlsafe(24)}
                atomic_write(recovery_file, json.dumps(recovery, indent=2))
            while True:
                try:
                    selector, candidate = prepare(page, host, b, entry, recovery[host], args.interactive)
                    try:
                        observe_browser(host, page.evaluate(FORM_STRUCTURE_JS))
                    except OSError:
                        print('Template evidence could not be saved; publishing safeguards remain active.', flush=True)
                    break
                except RefillRequested:
                    print('Refilling the approved details. No submission has been made.', flush=True)
            state.setdefault(host, {}).update({k: recovery[host][k] for k in ('omitted_fields', 'content_notes') if k in recovery[host]})
            atomic_write(recovery_file, json.dumps(recovery, indent=2))
            page.locator(selector).wait_for(state='visible')
            page.locator(selector).click(trial=True)
            state[host] = {**{k: recovery[host][k] for k in ('omitted_fields', 'content_notes') if k in recovery[host]}, 'status': 'submitting', 'reason': 'Submission started; do not automatically repeat', 'checked_utc': stamp()}
            if candidate:
                state[host]['candidate_url'] = candidate
            save(state, folder, entries)  # Journal before the irreversible click.
            started = True
            before = page.url
            check_cancelled()
            page.locator(selector).click()
            try:
                page.wait_for_url(lambda u: urlsplit(u)._replace(fragment='').geturl() != urlsplit(before)._replace(fragment='').geturl(), timeout=20000, wait_until='domcontentloaded')
            except BrowserTimeout:
                pass
            body = page.locator('body').inner_text()
            if host in CLASSIFIEDS and 'Your Ad successfully posted' in body:
                # Receipts can contain private edit links. Keep them only in the private recovery file.
                edit = re.search(r'https?://[^\s<>]+[?&]view=edit[^\s<>]+', body, re.I)
                if edit and domain(edit.group()) == host:
                    recovery[host]['edit_url'] = edit.group()
                    state[host]['submission_id'] = parse_qs(urlsplit(edit.group()).query).get('adid', [''])[0]
                    atomic_write(recovery_file, json.dumps(recovery, indent=2))
                go(page, recovery[host]['category_url'])
                links = page.locator('a[href]').filter(has_text=b['name']).evaluate_all('(els)=>els.map(e=>e.href)')
                links = list(dict.fromkeys(u for u in links if domain(u) == host))
                if len(links) == 1:
                    candidate = links[0]
                else:
                    state[host].update(status='pending', reason='Site accepted the ad; its public listing is not available yet')
            if state[host]['status'] == 'pending':
                pass
            elif 'The following errors were found when posting your ad' in body:
                state[host].pop('candidate_url', None)
                state[host].update(status='rejected', reason='Site rejected the ad form; inspect its validation or verification requirements')
            elif re.search(r'(pending (approval|review)|awaiting (approval|review)|submitted.{0,40}(approval|review)|reviewed by.{0,30}(editor|admin))', body, re.I):
                state[host].update(status='pending', reason='Submitted; awaiting site approval')
            else:
                if host == 'ikeanded.com':
                    if 'Thank you for running your listing' in body:
                        go(page, 'https://www.ikeanded.com/directory/myaccountlistings.aspx')
                        links = page.locator('a[href*="listingdisplay.aspx?lid="]').filter(has_text=b['name']).evaluate_all('(els)=>els.map(e=>e.href)')
                        new_links = [u for u in links if u not in recovery[host].get('existing_listing_urls', [])]
                        if len(new_links) == 1:
                            candidate = new_links[0]
                if not candidate and host != 'ikeanded.com' and urlsplit(page.url)._replace(fragment='').geturl() != urlsplit(before)._replace(fragment='').geturl() and domain(page.url) == host:
                    candidate = page.url
                if candidate:
                    state[host]['candidate_url'] = candidate
                    save(state, folder, entries)
                    try:
                        state[host] = {**state[host], **verify(browser, candidate, b, host)}
                    except Exception:
                        # A newly created page may take a moment to become publicly available.
                        time.sleep(2)
                        try:
                            state[host] = {**state[host], **verify(browser, candidate, b, host)}
                        except Exception:
                            state[host].update(status='unverified', reason='Submission attempted; candidate does not contain a verified public backlink')
                else:
                    state[host].update(status='uncertain', reason='Submission attempted; no public URL returned; do not resubmit automatically')
            if host in DIRECTORIES and state[host].get('status') in ('pending', 'rejected', 'uncertain'):
                recover_directory(browser, b, host, state, folder, entries,
                                  'existing' if state[host]['status'] == 'rejected' else 'new')
        except NeedsAttention as exc:
            deferred = not started and any(word in str(exc).lower() for word in ('captcha', 'challenge', 'verification', 'credentials', 'login email/password', 'skipped'))
            state[host] = {'status': 'uncertain' if started else 'skipped' if 'skipped by operator' in str(exc).lower() else 'needs_attention', 'score_deferred': deferred, 'reason': str(exc), 'checked_utc': stamp(), **({'candidate_url': state[host]['candidate_url']} if started and state[host].get('candidate_url') else {})}
        except Exception as exc:
            # Never persist raw browser exceptions: they can include filled passwords or page data.
            state.setdefault(host, {})
            state[host].update(status='uncertain' if started else 'failed', score_deferred=False, reason=type(exc).__name__ + ': network, changed form, or editor unavailable', checked_utc=stamp())
        finally:
            state.setdefault(host, {}).update({k: recovery.get(host, {})[k] for k in ('omitted_fields', 'content_notes') if k in recovery.get(host, {})})
            if started and not state.get(host, {}).get('candidate_url') and state.get(host, {}).get('status') not in ('published', 'rejected'):
                # Keep the result URL even if a later navigation or read raises an exception.
                current = urlsplit(page.url)
                try:
                    if domain(page.url) == host and current.path.rstrip('/') not in ('', '/submit', '/submit.php', '/login', '/signup') and page.url != before:
                        state[host]['candidate_url'] = page.url
                except ValueError:
                    pass
            if host == 'ikeanded.com' and state.get(host, {}).get('status') == 'published':
                state[host]['reason'] += '; free plan lasts 5 days'
            save(state, folder, entries)
            if review_interactive and state.get(host, {}).get('status') != 'published' and sys.exc_info()[0] is None:
                manual_review(browser, page, host, b, state, folder, entries, state[host].get('reason', 'Complete this site in the browser.'))
            if state.get(host, {}).get('status') != 'published':
                try:
                    diagnostic = page.locator('body').inner_text(timeout=2000)
                    for secret in (entry['password'], entry['username'], entry['login_email'], recovery.get(host, {}).get('edit_code', '')):
                        if secret:
                            diagnostic = diagnostic.replace(secret, '[REDACTED]')
                    diagnostic = re.sub(r'https?://[^\s<>]+[?&](?:view=edit|codemd5=|token=)[^\s<>]*', '[PRIVATE EDIT URL REDACTED]', diagnostic, flags=re.I)
                    (private / (host + '-last-page.txt')).write_text(diagnostic[:25000], encoding='utf-8')
                except Exception:
                    pass
            ctx.close()
        print(host, state[host]['status'], state[host].get('url', state[host].get('reason', '')), flush=True)
        if state[host].get('status') in ('uncertain', 'submitting', 'unverified'):
            raise InspectionRequired(host, 'Result remains unresolved after inspection; queue paused without reposting')
        time.sleep(1)


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('workbook', type=Path)
    parser.add_argument('--sheet')
    parser.add_argument('--campaign', default='default', help='Reuse to resume; change for another deliberate new campaign')
    parser.add_argument('--pool', type=Path, default=DEFAULT_POOL, help='Shared editable website pool')
    parser.add_argument('--workbook-sites-only', action='store_true', help='Use only the business workbook Sites instead of the shared pool')
    parser.add_argument('--target', type=int, default=DEFAULT_TARGET, help='Verified distinct websites, including existing results (default 100)')
    parser.add_argument('--only', nargs='+', help='Limit to these listing domains')
    parser.add_argument('--start-with', help='Process this workbook domain first, then continue with the rest')
    parser.add_argument('--audit', action='store_true', help='Inspect submission pages without publishing')
    parser.add_argument('--publish', action='store_true', help='Create new listings using implemented workflows')
    parser.add_argument('--interactive', action='store_true', help='Show browser, pause for CAPTCHA, and offer manual completion for unfinished sites')
    args = parser.parse_args()
    b, entries = load_workbook(args.workbook, args.sheet)
    if not args.workbook_sites_only:
        entries = merge_pool(entries, args.pool, args.workbook)
    if args.only:
        args.only = [domain(x) for x in args.only]
        if set(args.only) - {e['domain'] for e in entries}:
            raise ValueError('--only includes a domain absent or disabled in the selected website pool.')
    if args.start_with:
        args.start_with = domain(args.start_with)
        if args.start_with not in {e['domain'] for e in entries}:
            raise ValueError('--start-with includes a domain absent from the workbook.')
    folder = campaign_folder(b, args.campaign)
    folder.mkdir(parents=True, exist_ok=True)
    print(f'{len(entries)} rows; {len(set(e["domain"] for e in entries))} distinct domains. Output: {folder}', flush=True)
    with campaign_lock(folder / '.lock'):
        state_path = folder / 'state.json'
        state = json.loads(state_path.read_text(encoding='utf-8')) if state_path.exists() else {}
        configure_target(folder, args.target)
        seed_history(folder, state)
        save(state, folder, entries)
        if not args.audit and not args.publish:
            print('Workbook validated. Use --audit for inspection or --publish to submit.')
            return
        if args.publish and verified_count(state) >= args.target:
            print(progress_message(state, args.target))
            return
        if args.audit:
            audit([e for e in entries if not args.only or e['domain'] in args.only], folder)
        if not args.publish:
            return
        with sync_playwright() as p:
            browser = p.chromium.launch(headless=not args.interactive)
            try:
                publish(browser, b, entries, state, folder, args)
            finally:
                browser.close()
                if not args.workbook_sites_only:
                    try:
                        refresh_pool(args.pool, campaign=folder)
                    except PermissionError:
                        print('Close Website_Pool.xlsx and run python website_pool.py --refresh to update its scores.')
        print(progress_message(state, args.target), flush=True)
        print(f'Published URLs: {folder / "published_urls.txt"}', flush=True)


if __name__ == '__main__':
    try:
        main()
    except (NeedsAttention, ValueError, FileNotFoundError, KeyError) as exc:
        sys.exit(str(exc))
    except KeyboardInterrupt:
        sys.exit('Stopped. Resume with the same command; submitted jobs will not be blindly repeated.')
