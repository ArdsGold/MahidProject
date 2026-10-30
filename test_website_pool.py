"""Offline pool, evidence, target and API receipt checks. Never publishes live."""
import json
import tempfile
import unittest
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import patch

import openpyxl

from backlinks import publish, publish_telegraph, save
from uivision_publish import run_queue
from website_pool import (begin_attempt, configure_target, merge_pool, ranked_entries,
                          ratings, read_pool, refresh_pool, seed_history, sync_history,
                          verified_count)
from workbook_layout import create_template
from mappings.telegraph import payload


class PoolChecks(unittest.TestCase):
    def test_recorded_business_subdomain_counts_its_actual_hostname_once(self):
        from mappings.common import publication_domain
        record = {'status': 'published', 'publication_host': 'gym.blog.example',
                  'url': 'https://gym.blog.example/post'}
        self.assertEqual(publication_domain('blog.example', record), 'gym.blog.example')
        self.assertEqual(verified_count({'blog.example': record}), 1)
        self.assertEqual(verified_count({'blog.example': record, 'gym.blog.example': record}), 1)
        self.assertEqual(verified_count({'blog.example': {**record, 'publication_host': 'other.example'}}), 0)
        self.assertEqual(verified_count({'blog.example': {**record, 'url': 'https://other.blog.example/post'}}), 0)
        self.assertEqual(verified_count({'blog.example': {**record, 'publication_host': 'nested.gym.blog.example'}}), 0)

    def test_editable_rows_credentials_exclusions_and_duplicate_validation(self):
        with tempfile.TemporaryDirectory() as d:
            root = Path(d)
            path = root / 'pool.xlsx'
            refresh_pool(path, root)
            w = openpyxl.load_workbook(path)
            s = w['Websites']
            s.delete_rows(2, s.max_row)
            s.append(['https://www.rentry.co/', 'Yes', 'My notes'])
            s.append(['https://new.example/post', '', 'User added'])
            s.append(['disabled.example', 'No'])
            w.save(path)
            w.close()
            business = root / 'business.xlsx'
            create_template(business, ['rentry.co'], {'name': 'Gym', 'website': 'https://gym.example/', 'description': 'Approved'})
            entries = [{'row': 2, 'domain': 'rentry.co', 'username': 'private-user',
                        'password': 'secret', 'login_email': 'private@example.com'}]
            merged = merge_pool(entries, path, business)
            self.assertEqual([e['domain'] for e in merged], ['rentry.co', 'new.example'])
            self.assertEqual(merged[0]['password'], 'secret')
            self.assertEqual(merged[1]['password'], '')
            refresh_pool(path, root)
            self.assertEqual(read_pool(path)[0]['notes'], 'My notes')
            w = openpyxl.load_workbook(path)
            self.assertNotIn('secret', str(list(w['Websites'].values)))
            w['Websites'].append(['https://rentry.co/another', 'Yes'])
            w.save(path)
            w.close()
            with self.assertRaisesRegex(ValueError, 'Duplicate'):
                read_pool(path)
            w = openpyxl.load_workbook(path)
            w['Websites'].delete_rows(5)
            w['Websites']['A3'] = '=HYPERLINK("https://unsafe.example")'
            w.save(path)
            w.close()
            with self.assertRaisesRegex(ValueError, 'formula'):
                read_pool(path)

    def test_business_disabled_site_remains_excluded(self):
        with tempfile.TemporaryDirectory() as d:
            path, business = Path(d) / 'pool.xlsx', Path(d) / 'business.xlsx'
            refresh_pool(path, Path(d))
            create_template(business, ['rentry.co'])
            w = openpyxl.load_workbook(business)
            w['Sites']['E2'] = 'No'
            w.save(business)
            w.close()
            self.assertNotIn('rentry.co', [e['domain'] for e in merge_pool([], path, business)])

    def test_history_counts_attempts_once_and_combines_businesses(self):
        with tempfile.TemporaryDirectory() as d:
            root = Path(d)
            first, second = root / 'first', root / 'second'
            first.mkdir()
            second.mkdir()
            state = {'old.example': {'status': 'published', 'url': 'https://old.example/post'},
                     'deferred.example': {'status': 'needs_attention'},
                     'duplicate.example': {'status': 'rejected'},
                     'found.example': {'status': 'published', 'listing_origin': 'existing'}}
            seed_history(first, state)
            seed_history(first, state)
            begin_attempt(first, 'new.example')
            state['new.example'] = {'status': 'unverified'}
            sync_history(first, state)
            state['new.example'] = {'status': 'published', 'url': 'https://new.example/post'}
            for _ in range(3):
                sync_history(first, state)
                seed_history(first, state)
            s = ratings(root)
            self.assertEqual((s['new.example']['success'], s['new.example']['error']), (1, 0))
            self.assertEqual(s['new.example']['confidence'], 66.7)
            self.assertEqual(s['new.example']['historical'], 0)
            self.assertIsNone(s['deferred.example']['confidence'])
            self.assertIsNone(s['found.example']['confidence'])
            self.assertEqual(s['found.example']['recovered'], 1)
            begin_attempt(first, 'new.example')
            state['new.example'] = {'status': 'failed'}
            sync_history(first, state)
            begin_attempt(second, 'new.example')
            sync_history(second, {'new.example': {'status': 'published'}})
            s = ratings(root)
            self.assertEqual((s['new.example']['success'], s['new.example']['error']), (2, 1))
            self.assertEqual(s['new.example']['confidence'], 60)
            entries = [{'domain': h} for h in ['duplicate.example', 'untested.example', 'old.example']]
            with patch('website_pool.ROOT', root):
                ranked = ranked_entries(entries, first)
            self.assertEqual([e['domain'] for e in ranked], ['old.example', 'untested.example', 'duplicate.example'])

    def test_target_includes_existing_distinct_sites_and_stops_both_runners(self):
        state = {'old.example': {'status': 'published', 'url': 'https://old.example/post'},
                 'notpublic.example': {'status': 'pending', 'candidate_url': 'https://notpublic.example/post'},
                 'invalid.example': {'status': 'published', 'url': 'https://other.example/post'}}
        self.assertEqual(verified_count(state), 1)
        with tempfile.TemporaryDirectory() as d:
            folder = Path(d)
            configure_target(folder, 2)
            entries = [{'row': i, 'domain': h} for i, h in enumerate(['efdir.com', 'ifidir.com'])]
            def stage(client, job, action):
                if action == 'prepare':
                    return {'status': 'ready', 'challenge': False}
                self.assertEqual(json.loads((folder / 'state.json').read_text())[job['domain']]['status'], 'submitting')
                return {'status': 'candidate', 'url': 'https://' + job['domain'] + '/listing/one'}
            def verified(browser, url, business, host):
                return {'status': 'published', 'url': url}
            with patch('uivision_publish.run_stage', side_effect=stage) as stages, patch('uivision_publish.verify', side_effect=verified):
                run_queue(None, None, {'name': 'Gym', 'website': 'https://gym.example/', 'description': 'Approved'}, entries, state, folder)
                self.assertEqual(stages.call_count, 2)
            self.assertEqual(verified_count(state), 2)
            self.assertNotIn('ifidir.com', state)
            class NoBrowser:
                def new_context(self):
                    raise AssertionError('Target met; no browser should open')
            with patch('backlinks.ROOT', folder):
                publish(NoBrowser(), {}, entries, state, folder, SimpleNamespace(only=None, interactive=False))
            progress = json.loads((folder / 'progress.json').read_text())
            self.assertTrue(progress['target_met'])
            self.assertEqual(progress['remaining'], 0)
            with self.assertRaises(ValueError):
                configure_target(folder, 0)

    def test_error_continues_and_individual_platform_uses_existing_runner(self):
        with tempfile.TemporaryDirectory() as d:
            folder = Path(d)
            state = {}
            entries = [{'row': 1, 'domain': 'efdir.com'}, {'row': 2, 'domain': 'ifidir.com'}]
            def stage(client, job, action):
                if job['domain'] == 'efdir.com':
                    raise ValueError('Invalid changed form')
                return {'status': 'ready', 'challenge': False} if action == 'prepare' else {'status': 'pending'}
            with patch('uivision_publish.run_stage', side_effect=stage):
                run_queue(None, None, {'name': 'Gym', 'website': 'https://gym.example/', 'description': 'Approved'}, entries, state, folder)
            self.assertEqual(state['efdir.com']['status'], 'failed')
            self.assertEqual(state['ifidir.com']['status'], 'pending')
            self.assertEqual(ratings(folder.parent, folder)['efdir.com']['error'], 1)
            with patch('uivision_publish.publish') as delegated:
                run_queue(None, None, {}, [{'row': '', 'domain': 'telegra.ph'}], {}, folder)
                delegated.assert_called_once()
                self.assertEqual(delegated.call_args.args[-1].only, ['telegra.ph'])

    def test_telegraph_public_payload_limit_and_journal_before_api_publish(self):
        business = {'name': 'Example Gym', 'website': 'https://gym.example/location/',
                    'description': 'Approved ' + '\U0001f642' * 22000, 'hours': 'Monday: 15:00-21:00',
                    'keywords': 'gym, cheer', 'password': 'secret-password'}
        values, omitted = payload(business)
        self.assertLessEqual(len(values['content'].encode('utf-8')), 65536)
        self.assertIn('description (shortened)', omitted)
        self.assertIn('15:00-21:00', values['content'])
        self.assertIn('gym, cheer', values['content'])
        self.assertNotIn('secret-password', values['content'])
        with tempfile.TemporaryDirectory() as d:
            folder = Path(d)
            private = folder / 'private.json'
            state = {}
            def api(method, **kwargs):
                if method == 'createAccount':
                    return {'access_token': 'private-token'}
                journal = json.loads((folder / 'state.json').read_text())
                self.assertEqual(journal['telegra.ph']['status'], 'submitting')
                raise KeyboardInterrupt
            with patch('mappings.telegraph.call', side_effect=api):
                with self.assertRaises(KeyboardInterrupt):
                    publish_telegraph(None, business, [{'row': '', 'domain': 'telegra.ph'}], state, folder, {}, private)
            self.assertEqual(state['telegra.ph']['status'], 'submitting')
            self.assertNotIn('private-token', (folder / 'state.json').read_text())
            self.assertEqual(json.loads(private.read_text())['telegra.ph']['access_token'], 'private-token')


if __name__ == '__main__':
    unittest.main()
