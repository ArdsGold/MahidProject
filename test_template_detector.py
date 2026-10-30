"""Offline template recognition, privacy and pool integration. No live submissions."""
import json
import tempfile
import unittest
from collections import Counter
from pathlib import Path
from unittest.mock import patch

import openpyxl
from playwright.sync_api import sync_playwright

from template_detector import (FORM_STRUCTURE_JS, detect, detect_forms, load_catalog,
                               observe_browser, pool_labels, probe, save_catalog, scan)
from website_pool import HEADERS, read_pool, refresh_pool


FORM = '''<form method="post"><input name="TITLE" value="Private client name">
<input name="URL"><textarea name="DESCRIPTION">Private draft</textarea>
<input name="OWNER_NAME"><input name="OWNER_EMAIL"><select name="CATEGORY_ID"></select>
<input type="hidden" name="csrf" value="secret-token"><input type="submit" value="Continue"></form>'''


class TemplateChecks(unittest.TestCase):
    def test_browser_collector_matches_html_and_excludes_filled_values(self):
        with sync_playwright() as p:
            browser = p.chromium.launch()
            page = browser.new_page()
            page.set_content(FORM)
            snapshot = page.evaluate(FORM_STRUCTURE_JS)
            self.assertNotIn('Private', json.dumps(snapshot))
            self.assertNotIn('secret-token', json.dumps(snapshot))
            self.assertEqual(detect_forms(snapshot['forms'])['fingerprint'], detect(FORM)['fingerprint'])
            snapshot['host'] = 'new.test'
            with patch('template_detector.load_catalog', return_value={}), patch('template_detector.save_catalog') as saved:
                observe_browser('other.test', snapshot)
                saved.assert_not_called()
                observe_browser('new.test', snapshot)
                record = saved.call_args.args[0]['new.test']
                self.assertEqual(record['basis'], 'Live browser')
                self.assertEqual(record['family'], 'php-directory')
            browser.close()

    def test_structure_ignores_content_tokens_branding_and_detects_variants(self):
        first = detect(FORM)
        second = detect('<h1>Another website</h1>' + FORM.replace('secret-token', 'other-secret').replace('Private client name', 'Another client'))
        self.assertEqual(first['family'], 'php-directory')
        self.assertEqual(first['fingerprint'], second['fingerprint'])
        self.assertNotIn('secret', json.dumps(first))
        self.assertNotIn('Private', json.dumps(first))
        limited = detect(FORM.replace('name="DESCRIPTION"', 'name="DESCRIPTION" maxlength="500"'))
        self.assertEqual(limited['family'], 'php-directory')
        self.assertNotEqual(first['fingerprint'], limited['fingerprint'])
        self.assertEqual(detect(FORM.replace('method="post"', 'method="get"'))['family'], 'unmatched-form')
        self.assertEqual(detect(FORM.replace('name="TITLE"', 'name="TITLE" type="checkbox"'))['family'], 'unmatched-form')
        self.assertEqual(detect('<form method="post"></form>')['status'], 'No form observed')

    def test_unrelated_or_split_controls_do_not_match_and_challenges_stay_unknown(self):
        split = FORM.replace('<input name="OWNER_NAME">', '</form><form method="post"><input name="OWNER_NAME">')
        self.assertEqual(detect(split)['family'], 'unmatched-form')
        self.assertEqual(detect('<h1>PHP Directory</h1>')['family'], 'unknown')
        self.assertEqual(detect('<title>Just a moment</title>')['status'], 'Blocked / challenge')
        self.assertEqual(detect('<title>Just a moment</title><form><input name="verify"></form>')['status'], 'Blocked / challenge')
        classified = '<form method="post"><input id="adtitle"><textarea name="addesc"></textarea><input name="email"><select name="arealist"></select></form>'
        self.assertEqual(detect(classified)['family'], 'xzero-classifieds')
        self.assertEqual(detect(classified.replace('<select name="arealist"></select>', ''))['family'], 'unmatched-form')

    def test_outage_preserves_observed_structure_and_never_enables_workflow(self):
        records = {'unknown.test': {**detect(FORM), 'basis': 'Saved form'}}
        rows = [{'domain': 'unknown.test', 'enabled': 'Yes'}]
        with patch('template_detector.probe', return_value=('unknown.test', {'status': 'HTTP / network error', 'checked_utc': 'today'})):
            scan(rows, records, workers=1)
        fingerprint = records['unknown.test']['fingerprint']
        labels = pool_labels('unknown.test', records, Counter({fingerprint: 2}))
        self.assertEqual(labels[1], 'php-directory')
        self.assertEqual(labels[2], 'Shared structure (2 sites)')
        self.assertEqual(labels[5], 'HTTP / network error')
        self.assertEqual(labels[8], 'Needs mapping')
        self.assertEqual(pool_labels('not-seen.test', {}, Counter())[1], 'unknown')

    def test_pool_refresh_preserves_user_rows_and_campaign_confidence(self):
        with tempfile.TemporaryDirectory() as d:
            root = Path(d)
            path, catalog = root / 'pool.xlsx', root / 'catalog.json'
            w = openpyxl.Workbook()
            w.active.title = 'Websites'
            w.active.append(['Website', 'Enabled', 'Notes'])
            w.active.append(['https://one.test/post', 'Yes', 'Keep this note'])
            w.active.append(['two.test', 'No', 'Disabled'])
            w.save(path)
            w.close()
            records = {h: {**detect(FORM), 'basis': 'Saved form'} for h in ('one.test', 'two.test')}
            save_catalog(records, catalog)
            self.assertEqual(load_catalog(catalog), records)
            before = read_pool(path)
            with patch('template_detector.load_catalog', return_value=records), patch('website_pool.ratings', return_value={'one.test': {'confidence': 66.7}}):
                refresh_pool(path, root)
            self.assertEqual(read_pool(path), [dict(r, source='User added') for r in before])
            w = openpyxl.load_workbook(path)
            s = w['Websites']
            columns = {v: i for i, v in enumerate(HEADERS, 1)}
            self.assertEqual(s.cell(2, columns['Template family']).value, 'php-directory')
            self.assertEqual(s.cell(2, columns['Confidence %']).value, 66.7)
            self.assertEqual(s.cell(2, columns['Template automation']).value, 'Needs mapping')
            self.assertIn('Template Groups', w.sheetnames)
            self.assertEqual(s.auto_filter.ref, s.dimensions)
            w.close()

    def test_probe_only_gets_public_html_and_rejects_cross_domain_redirect(self):
        class Response:
            url = 'https://other.test/login?token=private'
            def __enter__(self): return self
            def __exit__(self, *args): pass
        row = {'domain': 'unknown.test', 'website': 'https://unknown.test/', 'enabled': 'Yes'}
        with patch('template_detector.urlopen', return_value=Response()) as request:
            host, result = probe(row)
            self.assertEqual(request.call_args.args[0].get_method(), 'GET')
            self.assertEqual(result['status'], 'HTTP / network error')
            self.assertNotIn('token', json.dumps(result))


if __name__ == '__main__':
    unittest.main()
