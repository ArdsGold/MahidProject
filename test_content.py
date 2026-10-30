"""Public-content fitting and omission reporting; synthetic data, no publishing."""
import csv
import json
import tempfile
import unittest
import openpyxl
from pathlib import Path
from unittest.mock import patch

from playwright.sync_api import sync_playwright

from mappings.common import CONTENT_FITTER, NeedsAttention, content, content_parts, fit_body
from mappings.editors import article_html
from uivision_publish import make_job, run_queue
from backlinks import load_workbook


BUSINESS = {'name': 'Example Gym', 'website': 'https://business.example/',
            'address': '10 Main St, Example City', 'phone': '+1 555 0100',
            'email': 'public@example.com', 'description': 'Approved description. ' * 20,
            'category': 'Sports', 'hours': 'Monday: 09:00–17:00', 'keywords': 'Cheer, Gym',
            'facebook': 'https://facebook.com/example', 'instagram': 'https://instagram.com/example',
            'logo': 'https://business.example/logo.png', 'timezone': 'America/Chicago',
            'password': 'SECRET-PASSWORD', 'login_email': 'SECRET-LOGIN', 'edit_code': 'SECRET-EDIT'}


class ContentChecks(unittest.TestCase):
    def test_legacy_hour_range_separator_uses_confirmed_hyphen(self):
        with tempfile.TemporaryDirectory() as d:
            path = Path(d) / 'client.xlsx'
            workbook = openpyxl.Workbook()
            sheet = workbook.active
            for row in [('Business Name', 'Example Gym'), ('Website', 'https://business.example/'),
                        ('Description', 'Approved'), ('Business Hours', 'Monday 3\ufffd9 PM\nTuesday 12\u20139 PM'),
                        ('Instagram', 'https://instagram.com/example'), ('Listing Sites', ''), ('pastelink.net', '')]:
                sheet.append(row)
            workbook.save(path)
            workbook.close()
            business, _ = load_workbook(path)
            self.assertEqual(business['hours'], 'Monday 3-9 PM\nTuesday 12-9 PM')
            plain, _ = content(business)
            self.assertIn('Hours: Monday 3-9 PM', plain)
            self.assertIn('Instagram: https://instagram.com/example', plain)

    def test_all_public_fields_in_plain_markdown_html_and_browser_payload(self):
        plain, markdown = content(BUSINESS)
        html = article_html(BUSINESS)
        job = make_job(BUSINESS, 'foldads.com')
        for value in (plain, markdown, html, json.dumps(job, ensure_ascii=False)):
            for key in ('name', 'website', 'address', 'phone', 'email', 'category', 'hours',
                        'keywords', 'facebook', 'instagram', 'logo', 'timezone'):
                self.assertIn(BUSINESS[key], value)
            self.assertNotIn('SECRET-', value)
        escaped = article_html({**BUSINESS, 'description': '<script>not executable</script>'})
        self.assertIn('&lt;script&gt;', escaped)
        self.assertNotIn('<script>', escaped)

    def test_native_limits_preserve_essentials_report_shortening_and_block_too_small(self):
        with sync_playwright() as p:
            browser = p.chromium.launch()
            page = browser.new_page()
            page.set_content('<textarea id="body" maxlength="230"></textarea>')
            record = {}
            fitted = fit_body(page, '#body', BUSINESS, record)
            self.assertLessEqual(len(fitted.encode('utf-16-le')) // 2, 230)
            for key in ('name', 'website', 'address', 'phone', 'email'):
                self.assertIn(BUSINESS[key], fitted)
            self.assertIn('description (shortened)', record['omitted_fields'])
            self.assertIn('keywords (character limit)', record['omitted_fields'])
            self.assertEqual(page.locator('#body').input_value(), fitted)
            page.locator('#body').evaluate('(e)=>e.maxLength=10')
            with self.assertRaises(NeedsAttention):
                fit_body(page, '#body', BUSINESS, record)
            self.assertIn('website (character limit)', record['omitted_fields'])
            # Emoji must use the browser's two-unit UTF-16 count.
            parts = content_parts({**BUSINESS, 'name': 'Gym 🏆'})
            for limit in range(0, 260):
                result = page.evaluate(CONTENT_FITTER, {**parts, 'limit': limit})
                self.assertLessEqual(len(result['text'].encode('utf-16-le')) // 2, limit)
            browser.close()

    def test_damaged_hours_are_reported_and_structured_address_is_included(self):
        parts = content_parts({**BUSINESS, 'hours': 'Monday 3\ufffd9 PM'})
        self.assertIn('hours (damaged source text)', parts['omitted_fields'])
        self.assertNotIn('\ufffd', '\n'.join(p['text'] for p in parts['parts']))
        plain, _ = content({'name': 'Example', 'website': 'https://example.com/', 'description': 'Approved',
                            'street': '1 Main', 'address_line_2': 'Suite 2', 'city': 'Town'})
        self.assertIn('Address: 1 Main, Suite 2, Town', plain)

    def test_omissions_survive_journal_submission_verification_and_csv(self):
        host = 'smartdir.org'
        entries = [{'row': 1, 'domain': host}]
        metadata = {'omitted_fields': ['keywords (character limit)'], 'content_notes': 'Combined text limit: 230 UTF-16 units'}
        with tempfile.TemporaryDirectory() as d:
            folder, state = Path(d), {}
            def stage(client, job, action):
                if action == 'prepare':
                    return {'status': 'ready', 'challenge': False, 'content_limit': 230, **metadata}
                self.assertEqual(job['content_limit'], 230)
                journal = json.loads((folder / 'state.json').read_text())
                self.assertEqual(journal[host]['omitted_fields'], metadata['omitted_fields'])
                return {'status': 'candidate', 'url': 'https://smartdir.org/details.php?id=42'}
            with patch('uivision_publish.run_stage', side_effect=stage), patch('uivision_publish.verify',
                       return_value={'status': 'published', 'url': 'https://smartdir.org/details.php?id=42'}):
                run_queue(None, None, BUSINESS, entries, state, folder)
            self.assertEqual(state[host]['omitted_fields'], metadata['omitted_fields'])
            with (folder / 'results.csv').open(encoding='utf-8-sig', newline='') as f:
                row = next(csv.DictReader(f))
            self.assertEqual(row['status'], 'published')
            self.assertEqual(row['omitted_fields'], 'keywords (character limit)')
            self.assertEqual(row['content_notes'], metadata['content_notes'])


if __name__ == '__main__':
    unittest.main()
