"""Create and read the standard local publisher workbook. No network or publishing."""
import argparse
import re
from datetime import datetime, time
from pathlib import Path
from urllib.parse import urlsplit

import openpyxl
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.table import Table, TableStyleInfo


# Label, internal field, importance, instructions. Labels are matched by name, not row number.
FIELDS = [
    ('Business name', 'name', 'Required', 'Official public business name.'),
    ('Website', 'website', 'Required', 'Full destination URL, including https://.'),
    ('Description', 'description', 'Required', 'Approved business description. Plain text; line breaks are allowed.'),
    ('Short description', 'short_description', 'Optional', 'Approved summary, at most 160 characters. Otherwise the tool shortens the first sentence.'),
    ('Street address', 'street', 'Recommended', 'Street number and street name; do not include city or postal code.'),
    ('Address line 2', 'address_line_2', 'Optional', 'Suite, unit or floor, for example Ste 400.'),
    ('City', 'city', 'Recommended', 'City only, for example St. Louis.'),
    ('State / Region', 'region', 'Recommended', 'For US businesses use the two-letter code, for example MO.'),
    ('Postal code', 'postal_code', 'Recommended', 'Text format preserves leading zeros.'),
    ('Country', 'country', 'Recommended', 'Country name or code, for example United States or US.'),
    ('Phone', 'phone', 'Recommended', 'Public business number, including country code. Keep as text.'),
    ('Business email', 'email', 'Recommended', 'Public business email. Account login emails belong on Sites.'),
    ('Category', 'category', 'Recommended', 'Main business category. A site may require a different category name.'),
    ('Keywords', 'keywords', 'Optional', 'Comma-separated approved keywords.'),
    ('Time zone', 'timezone', 'Recommended', 'Local business time zone, for example America/Chicago. Hours are entered on Hours.'),
    ('Logo URL', 'logo', 'Optional', 'Public https:// URL to the business logo.'),
    ('Facebook URL', 'facebook', 'Optional', 'Public business profile URL.'),
    ('Instagram URL', 'instagram', 'Optional', 'Public business profile URL.'),
    ('LinkedIn URL', 'linkedin', 'Optional', 'Public business profile URL.'),
    ('YouTube URL', 'youtube', 'Optional', 'Public business channel URL.'),
    ('TikTok URL', 'tiktok', 'Optional', 'Public business profile URL.'),
    ('X / Twitter URL', 'twitter', 'Optional', 'Public business profile URL.'),
]
DAYS = ('Monday', 'Tuesday', 'Wednesday', 'Thursday', 'Friday', 'Saturday', 'Sunday')
HOURS_HEADERS = ('Day', 'Status', 'Opens', 'Closes')
SITE_HEADERS = ('Listing site', 'Username', 'Password', 'Login email', 'Enabled')


def cell_text(cell, location, strip=True):
    if cell.data_type == 'f':
        raise ValueError(location + ': use a value, not a formula.')
    value = str(cell.value) if cell.value is not None else ''
    return value.strip() if strip else value


def clock(value, location):
    if isinstance(value, datetime):
        value = value.time()
    if isinstance(value, time):
        if value.second or value.microsecond:
            raise ValueError(location + ': use hours and minutes only.')
        return value.strftime('%H:%M')
    if isinstance(value, (float, int)) and not isinstance(value, bool) and 0 <= value < 1:
        minutes = round(value * 1440)
        if minutes < 1440:
            return f'{minutes // 60:02}:{minutes % 60:02}'
    value = str(value).strip()
    if not re.fullmatch(r'(?:[01]\d|2[0-3]):[0-5]\d', value):
        raise ValueError(location + ': use 24-hour HH:MM, for example 15:00.')
    return value


def template_hours(value):
    """Convert approved source hours to separate ASCII time cells; never guess missing days."""
    result = {}
    for line in str(value or '').splitlines():
        if not line.strip():
            continue
        match = re.fullmatch(r'(' + '|'.join(DAYS) + r')\s*:?\s+(.+)', line.strip(), re.I)
        if not match:
            raise ValueError('Source hours: use a weekday followed by its hours: ' + line)
        day, hours = match.groups()
        day = day.title()
        if day in result:
            raise ValueError('Source hours: duplicate ' + day + '.')
        if hours.lower() in ('closed', '24 hours', 'by appointment'):
            result[day] = (hours.capitalize(), '', '')
            continue
        hours = re.sub(r'[\u2013\u2014\ufffd]', '-', hours).strip()
        times = re.fullmatch(r'(\d{2}:\d{2})\s*-\s*(\d{2}:\d{2})(?: \(closes next day\))?', hours)
        if times:
            start, end = (clock(v, 'Source hours: ' + day) for v in times.groups())
        else:
            times = re.fullmatch(r'(\d{1,2}(?::\d{2})?)\s*(AM|PM)?\s*-\s*(\d{1,2}(?::\d{2})?)\s*(AM|PM)', hours, re.I)
            if not times:
                raise ValueError('Source hours for ' + day + ': use HH:MM-HH:MM or explicit AM/PM times.')
            opening, ampm, closing, end_ampm = times.groups()
            def to_24h(text, suffix):
                text = text if ':' in text else text + ':00'
                return datetime.strptime(text + suffix.upper(), '%I:%M%p').strftime('%H:%M')
            start, end = to_24h(opening, ampm or end_ampm), to_24h(closing, end_ampm)
            if not ampm and start >= end:
                raise ValueError('Source hours for ' + day + ': specify AM/PM for both opening and closing times.')
        if start == end:
            raise ValueError('Source hours for ' + day + ': choose 24 hours instead of equal opening and closing times.')
        result[day] = ('Open', start, end)
    return result


def read_standard(w, normalize_domain):
    """Read the named sheets even when rows/columns move or another tab is active."""
    s = w['Business']
    labels = {label.lower(): key for label, key, _, _ in FIELDS}
    b = {}
    for row in s.iter_rows(min_row=2):
        label = cell_text(row[0], f'Business!A{row[0].row}').lower()
        if not label:
            continue
        if label not in labels:
            raise ValueError(f'Business!A{row[0].row}: unknown field label; use the template labels.')
        key = labels[label]
        if key in b:
            raise ValueError(f'Business: duplicate field {label}.')
        b[key] = cell_text(row[1], f'Business!B{row[1].row}')
    for key in ('name', 'website', 'description'):
        if not b.get(key):
            raise ValueError('Business: missing required field ' + key + '.')
    if urlsplit(b['website']).scheme not in ('http', 'https'):
        raise ValueError('Business Website must include http:// or https://.')
    normalize_domain(b['website'])
    if len(b.get('short_description', '')) > 160:
        raise ValueError('Business Short description must be at most 160 characters.')
    for key in ('logo', 'facebook', 'instagram', 'linkedin', 'youtube', 'tiktok', 'twitter'):
        if b.get(key):
            if urlsplit(b[key]).scheme not in ('http', 'https'):
                raise ValueError('Business: ' + key + ' must be a full HTTP(S) URL.')
            normalize_domain(b[key])
    street = ' '.join(filter(None, [b.get('street'), b.get('address_line_2')]))
    region_postal = ' '.join(filter(None, [b.get('region'), b.get('postal_code')]))
    b['address'] = ', '.join(filter(None, [street, b.get('city'), region_postal, b.get('country')]))

    if 'Hours' in w.sheetnames:
        h = w['Hours']
        headers = {cell_text(c, f'Hours!{c.coordinate}').lower(): c.column - 1 for c in h[1] if c.value is not None}
        if len(headers) != sum(c.value is not None for c in h[1]):
            raise ValueError('Hours: duplicate column headers.')
        if set(key.lower() for key in HOURS_HEADERS) - headers.keys():
            raise ValueError('Hours: keep Day, Status, Opens and Closes headers in row 1.')
        hours = {}
        for row in h.iter_rows(min_row=2):
            cells = {key.lower(): row[headers[key.lower()]] for key in HOURS_HEADERS}
            if not any(c.value not in (None, '') for c in cells.values()):
                continue
            day = cell_text(cells['day'], f'Hours!{cells["day"].coordinate}').title()
            if day not in DAYS or day in hours:
                raise ValueError(f'Hours!{cells["day"].coordinate}: use each Monday-Sunday day only once.')
            status = cell_text(cells['status'], f'Hours!{cells["status"].coordinate}').lower()
            for key in ('opens', 'closes'):
                cell_text(cells[key], f'Hours!{cells[key].coordinate}')  # Reject formulas, including time cells.
            opening, closing = cells['opens'].value, cells['closes'].value
            opening = None if opening == '' else opening
            closing = None if closing == '' else closing
            if not status:
                if opening is not None or closing is not None:
                    raise ValueError(f'Hours row {row[0].row}: select a status before entering times.')
                hours[day] = ''
            elif status == 'open':
                start = clock(opening, f'Hours!{cells["opens"].coordinate}')
                end = clock(closing, f'Hours!{cells["closes"].coordinate}')
                if start == end:
                    raise ValueError(f'Hours row {row[0].row}: equal times are ambiguous; choose 24 hours if appropriate.')
                hours[day] = start + '-' + end + (' (closes next day)' if end < start else '')
            elif status in ('closed', '24 hours', 'by appointment'):
                if opening is not None or closing is not None:
                    raise ValueError(f'Hours row {row[0].row}: clear Opens/Closes when status is not Open.')
                hours[day] = status.capitalize()
            else:
                raise ValueError(f'Hours!{cells["status"].coordinate}: choose Open, Closed, 24 hours, or By appointment.')
        b['hours'] = '\n'.join(day + ': ' + hours[day] for day in DAYS if hours.get(day))

    s = w['Sites']
    headers = {cell_text(c, c.coordinate).lower(): c.column - 1 for c in s[1] if c.value is not None}
    if len(headers) != len([c for c in s[1] if c.value is not None]):
        raise ValueError('Sites: duplicate column headers.')
    if set(h.lower() for h in SITE_HEADERS) - headers.keys():
        raise ValueError('Sites: keep all five template column headers in row 1.')
    entries = []
    for row in s.iter_rows(min_row=2):
        values = {name: cell_text(row[index], f'Sites!{row[index].coordinate}', strip=name != 'password') for name, index in headers.items()}
        site = values['listing site']
        if not site:
            if any(values.get(k) for k in ('username', 'password', 'login email')):
                raise ValueError(f'Sites row {row[0].row}: account details need a Listing site.')
            continue
        enabled = values['enabled'].lower()
        if enabled not in ('yes', 'no'):
            raise ValueError(f'Sites row {row[0].row}: Enabled must be Yes or No.')
        if enabled == 'no':
            continue
        entries.append({'row': row[0].row, 'domain': normalize_domain(site), 'username': values['username'],
                        'password': values['password'], 'login_email': values['login email']})
    if not entries:
        raise ValueError('Sites: enter at least one listing site with Enabled = Yes.')
    return b, entries


def create_template(path, sites, business=None):
    b = business or {}
    hours = template_hours(b.get('hours'))
    w = openpyxl.Workbook()
    s = w.active
    s.title = 'Business'
    s.append(['Field', 'Value - fill in blue cells', 'Importance', 'Guidance'])
    for label, key, required, guide in FIELDS:
        s.append([label, b.get(key, ''), required, guide])
        s.cell(s.max_row, 2).comment = Comment(guide, 'Local publisher template')
        s.row_dimensions[s.max_row].height = 32 if key != 'description' else 110
    s.column_dimensions['A'].width = 24
    s.column_dimensions['B'].width = 76
    s.column_dimensions['C'].width = 17
    s.column_dimensions['D'].width = 81
    s.freeze_panes = 'B2'

    h = w.create_sheet('Hours')
    h.append([*HOURS_HEADERS, 'Guidance'])
    for day in DAYS:
        h.append([day, *hours.get(day, ('', '', '')), 'Open: use 24-hour times, e.g. 15:00 and 21:00. Otherwise leave times empty. Blank status means unknown.'])
    for col, width in [('A', 18), ('B', 23), ('C', 15), ('D', 15), ('E', 100)]:
        h.column_dimensions[col].width = width
    h.freeze_panes = 'B2'
    status = DataValidation(type='list', formula1='"Open,Closed,24 hours,By appointment"', allow_blank=True)
    status.errorTitle, status.error = 'Choose a status', 'Select a value from the dropdown.'
    status.errorStyle = 'stop'
    status.promptTitle, status.prompt = 'Daily status', 'Select Open, Closed, 24 hours or By appointment. Blank means unknown.'
    status.showInputMessage = True
    status.showErrorMessage = True
    h.add_data_validation(status)
    status.add('B2:B8')
    # Separate time cells avoid all pasted range separators. Accept native Excel times too.
    time_rule = 'OR(C2="",AND(ISNUMBER(C2),C2>=0,C2<1),IFERROR(AND(LEN(C2)=5,TEXT(TIMEVALUE(C2),"hh:mm")=C2),FALSE))'
    valid_time = DataValidation(type='custom', formula1='AND($B2="Open",' + time_rule + ')', allow_blank=True)
    valid_time.errorTitle, valid_time.error = 'Use a 24-hour time', 'Select Open first. Enter HH:MM (15:00), not a range (3-9 PM).'
    valid_time.errorStyle, valid_time.showErrorMessage = 'stop', True
    valid_time.promptTitle, valid_time.prompt = 'Separate 24-hour time', 'Use 00:00-23:59, for example 15:00. Keep Closed/24 hours/By appointment times empty.'
    valid_time.showInputMessage = True
    h.add_data_validation(valid_time)
    valid_time.add('C2:D8')
    for row in range(2, 9):
        for col in (3, 4):
            h.cell(row, col).comment = Comment('24-hour HH:MM, for example 09:00 or 17:30. For overnight hours enter the actual closing time; the reader marks it as next day.', 'Template')

    t = w.create_sheet('Sites')
    t.append(SITE_HEADERS)
    for host in dict.fromkeys(sites):
        t.append([host, '', '', '', 'Yes'])
    if t.max_row == 1:
        t.append(['', '', '', '', 'Yes'])
    for col, width in [('A', 37), ('B', 30), ('C', 30), ('D', 42), ('E', 16)]:
        t.column_dimensions[col].width = width
    t.freeze_panes = 'B2'
    table = Table(displayName='ListingAccounts', ref=f'A1:E{t.max_row}')
    table.tableStyleInfo = TableStyleInfo(name='TableStyleMedium2', showRowStripes=True)
    t.add_table(table)
    enabled = DataValidation(type='list', formula1='"Yes,No"')
    enabled.showErrorMessage = True
    enabled.error = 'Choose Yes or No.'
    t.add_data_validation(enabled)
    enabled.add('E2:E10001')

    notes = w.create_sheet('Instructions')
    notes.append(['Topic', 'Instructions'])
    for topic, instruction in [
        ('One business', 'Save a new copy of this workbook for each business. Keep Business, Hours and Sites sheet names.'),
        ('Business', 'Fill blue value cells in column B. Required: Business name, Website and Description. Other fields may be required by individual sites.'),
        ('Address', 'Use separate Street address, Address line 2, City, State / Region, Postal code and Country fields. The tool assembles the full address.'),
        ('Text fields', 'Phone numbers and postal codes are text to preserve + signs and leading zeros. Paste values only, not formulas.'),
        ('Hours', 'Use one row per day. Open requires separate 24-hour HH:MM times, for example Opens 15:00 and Closes 21:00. Do not paste a range or dash into a time cell. Closed, 24 hours and By appointment require empty times. Blank means unknown, not closed.'),
        ('Overnight hours', 'Enter the actual closing time, for example Opens 22:00 and Closes 02:00. The reader marks this as closing the next day. Equal times require the 24 hours status.'),
        ('Time zone', 'Enter the business local time zone on Business; Hours are local wall-clock times and are not converted.'),
        ('Sites', 'The original workbook domains are prefilled. Credentials are blank. Add rows as needed. Enabled = Yes includes a site; No excludes it.'),
        ('Login vs public email', 'Business email is public. Login email on Sites is only for signing into that site. Fill credentials only when required.'),
        ('Keep labels', 'Keep Business field labels and Hours/Sites column headers. The reader matches Business labels and Hours/Sites headers by name; changing a header is an error.'),
        ('Blank vs example', 'The blank template has no business details or hours. The example includes supplied public details and hours converted into separate time cells. Credentials remain blank.'),
        ('What is used', 'Long description/body fields combine available public details, including hours, socials and keywords. Space limits may cause omissions reported in results.csv. Login credentials are excluded from public content.'),
        ('Checking pasted data', 'Excel validation guides manual entry. Paste/import can bypass it; run the validate-only command before publishing. Errors identify the sheet and cell to correct. Generated hours use standard ASCII hyphens.'),
        ('Validate only', 'Run: python backlinks.py "your-file.xlsx". Without --publish or --audit no site is accessed.'),
        ('Publish', 'Run: python backlinks.py "your-file.xlsx" --publish. Existing campaign results are resumed; template format does not add website coverage.'),
        ('Results', 'Published URLs and status reports are written under output/. The input workbook is never overwritten by the publisher.'),
        ('Private workbook', 'After adding passwords, keep this workbook private. The template contains no account credentials.'),
    ]:
        notes.append([topic, instruction])
    notes.column_dimensions['A'].width, notes.column_dimensions['B'].width = 27, 118
    notes.freeze_panes = 'B2'

    for sh in w:
        sh.sheet_view.showGridLines = False
        sh.sheet_properties.pageSetUpPr.fitToPage = True
        sh.page_setup.orientation = 'landscape'
        sh.page_setup.paperSize = sh.PAPERSIZE_A4
        sh.page_setup.fitToWidth, sh.page_setup.fitToHeight = 1, 0
        sh.print_title_rows = '1:1'
        sh.sheet_properties.tabColor = '1E5A86'
        for c in sh[1]:
            c.fill = PatternFill('solid', fgColor='173B55')
            c.font = Font(name='Calibri', size=11, color='FFFFFF', bold=True)
            c.alignment = Alignment(vertical='center', wrap_text=True)
        sh.row_dimensions[1].height = 32
        for row in sh.iter_rows(min_row=2):
            if not sh.row_dimensions[row[0].row].height:
                sh.row_dimensions[row[0].row].height = 40 if sh.title == 'Instructions' else 31
            for c in row:
                editable = sh.title == 'Business' and c.column == 2 or sh.title == 'Hours' and c.column in (2, 3, 4) or sh.title == 'Sites'
                c.number_format = '@'
                c.fill = PatternFill('solid', fgColor='EAF4FF' if editable else 'F3F5F7' if c.row % 2 == 0 else 'FFFFFF')
                c.font = Font(name='Calibri', size=11, color='174A7E' if editable else '344054')
                c.alignment = Alignment(vertical='top', wrap_text=True)
    w.save(path)
    w.close()


if __name__ == '__main__':
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--source', type=Path, help='Prefill source domains and create a public business example; never copy credentials')
    parser.add_argument('--output', type=Path, default=Path('templates'))
    args = parser.parse_args()
    sites, example = [], {}
    if args.source:
        from backlinks import load_workbook
        example, entries = load_workbook(args.source)
        sites = [e['domain'] for e in entries]
        match = re.fullmatch(r'(.+?)(?:\s+(Ste\s+\S+))?,\s*([^,]+),\s*([A-Z]{2})\s+(\S+),\s*(.+)', example.get('address', ''))
        if match:
            example.update(dict(zip(('street', 'address_line_2', 'city', 'region', 'postal_code', 'country'), match.groups(default=''))))
    args.output.mkdir(parents=True, exist_ok=True)
    create_template(args.output / 'Backlink_Template.xlsx', sites)
    if example:
        create_template(args.output / 'Cheer_Athletics_Example.xlsx', sites, example)
    print('Created template files in ' + str(args.output.resolve()) + '; no credentials copied.')
