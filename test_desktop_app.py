"""Offline desktop UI, operator protocol, safe pause and public-report export checks."""
import csv
import io
import json
import tempfile
import queue
import tkinter as tk
import unittest
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import Mock, patch

import openpyxl
import operator_notifications as notifications
from backlink_app import BacklinkApp, Worker, export_results, inspect_client
from desktop_runner import DesktopChannel, resolve_review, publish_from_app, connect_chrome
from mappings.common import InspectionRequired, NeedsAttention
from workbook_layout import create_template

BUSINESS = {'name': 'Example Gym', 'website': 'https://business.example/location/',
            'description': 'Approved public description'}


def fixtures(folder):
    client, pool = folder / 'client.xlsx', folder / 'pool.xlsx'
    create_template(client, ['rentry.co'], BUSINESS)
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = 'Websites'
    sheet.append(['Website', 'Enabled', 'Notes'])
    sheet.append(['https://rentry.co', 'Yes', ''])
    sheet.append(['https://unmapped.example', 'Yes', ''])
    workbook.save(pool)
    workbook.close()
    return client, pool


class DesktopChecks(unittest.TestCase):
    def tearDown(self):
        notifications.STOP_REQUESTED.clear()

    def test_connection_wakes_chrome_for_both_bridge_status_wordings(self):
        for status in ('No tool connection; wake channel chrome', 'NOT connected; wake channel chrome'):
            with self.subTest(status=status), patch('uivision.UiVision') as bridge:
                client = bridge.return_value.__enter__.return_value
                client.call.side_effect = [
                    {'content': [{'type': 'text', 'text': status}]},
                    {'content': [{'type': 'text', 'text': 'Panel opened'}]},
                    {'content': [{'type': 'text', 'text': 'Connected chrome#1'}]}]
                connect_chrome()
                self.assertEqual([c.args[0] for c in client.call.call_args_list], ['bridge_status', 'open_panel', 'bridge_status'])

    def test_connected_bridge_with_historical_hello_note_is_not_a_failure(self):
        with patch('uivision.UiVision') as bridge:
            client = bridge.return_value.__enter__.return_value
            client.call.return_value = {'content': [{'type': 'text', 'text':
                'Bridge v1.7.20: 2 browser connections:\n chrome#1: host not connected at hello time <-- ACTIVE\nWake channels: chrome'}]}
            connect_chrome()
            client.call.assert_called_once_with('bridge_status')

    def test_real_worker_pipe_delivers_prompt_response_and_completion(self):
        with tempfile.TemporaryDirectory() as d:
            folder = Path(d)
            (folder / 'desktop_runner.py').write_text('''import sys,json
print(json.dumps({'event':'action','id':'fixture','message':'Complete verification'}),flush=True)
answer=json.loads(sys.stdin.readline())
assert answer=={'command':'respond','id':'fixture','value':''}
print(json.dumps({'event':'finished','message':'Saved'}),flush=True)
''', encoding='utf-8')
            events = queue.Queue()
            with patch('backlink_app.ROOT', folder):
                worker = Worker([], events)
            self.assertEqual(events.get(timeout=10)['event'], 'action')
            worker.send('respond', id='fixture', value='')
            self.assertEqual(events.get(timeout=10)['event'], 'finished')
            self.assertEqual(events.get(timeout=10), {'event': 'exit', 'code': 0})

    def test_workbook_selection_keeps_campaign_identity_and_never_changes_journal(self):
        import backlinks
        with tempfile.TemporaryDirectory() as d, patch.object(backlinks, 'ROOT', Path(d)):
            workbook, pool = fixtures(Path(d))
            first = inspect_client(workbook, pool)
            first['folder'].mkdir(parents=True)
            state = {'rentry.co': {'status': 'submitting', 'reason': 'Protected'}}
            (first['folder'] / 'state.json').write_text(json.dumps(state))
            second = inspect_client(workbook, pool)
            self.assertEqual(first['folder'], second['folder'])
            self.assertEqual(json.loads((first['folder'] / 'state.json').read_text()), state)
            self.assertEqual(first['ready'], 1)
            self.assertEqual(len(first['hosts']), 2)

    def test_excel_export_includes_errors_and_verified_links_but_no_private_files(self):
        with tempfile.TemporaryDirectory() as d:
            folder = Path(d)
            with (folder / 'results.csv').open('w', newline='', encoding='utf-8-sig') as f:
                writer = csv.writer(f)
                writer.writerow(['domain', 'status', 'published_url', 'reason'])
                writer.writerow(['rentry.co', 'published', 'https://rentry.co/public', 'Verified'])
                writer.writerow(['unmapped.example', 'error', '', '=UNTRUSTED()'])
            (folder / 'published_urls.txt').write_text('https://rentry.co/public\n')
            (folder / 'recovery.json').write_text('{"password":"DO-NOT-EXPORT"}')
            destination = folder / 'export.xlsx'
            export_results(folder, destination)
            result = openpyxl.load_workbook(destination)
            self.assertEqual(result['Verified Links']['A2'].value, 'https://rentry.co/public')
            self.assertEqual(result['Results']['D3'].data_type, 's')
            self.assertNotIn('DO-NOT-EXPORT', str([list(s.values) for s in result]))
            self.assertEqual(result['Results']['B3'].value, 'error')
            result.close()

    def test_operator_response_requires_matching_request_and_notifies_once(self):
        channel = DesktopChannel(outgoing=io.StringIO())
        channel.answers.put({'id': 'old', 'value': 'skip'})
        channel.answers.put({'id': 'current', 'value': ''})
        with patch('desktop_runner.uuid.uuid4', return_value=SimpleNamespace(hex='current')), \
             patch.object(notifications, 'notify') as alert:
            self.assertEqual(channel.ask('Complete CAPTCHA'), '')
            alert.assert_called_once()
        event = json.loads(channel.outgoing.getvalue())
        self.assertEqual(event['event'], 'action')
        self.assertEqual(event['id'], 'current')

    def test_pause_interrupts_a_wait_without_answering_or_submitting(self):
        channel = DesktopChannel(outgoing=io.StringIO())
        notifications.STOP_REQUESTED.set()
        with patch.object(notifications, 'notify') as alert, self.assertRaises(KeyboardInterrupt):
            channel.ask('Complete CAPTCHA')
        alert.assert_not_called()
        from site_access import submit_account
        client = Mock()
        with self.assertRaises(KeyboardInterrupt):
            submit_account(client, {}, {}, {})
        client.call.assert_not_called()

    def test_pause_after_filling_never_starts_a_final_submission(self):
        from uivision_publish import run_queue
        with tempfile.TemporaryDirectory() as d:
            state = {}
            def stage(client, job, action):
                self.assertEqual(action, 'prepare')
                notifications.STOP_REQUESTED.set()
                return {'status': 'ready', 'challenge': False}
            with patch('uivision_publish.run_stage', side_effect=stage) as calls:
                with self.assertRaises(KeyboardInterrupt):
                    run_queue(None, None, BUSINESS, [{'row': 1, 'domain': 'getadsonline.com'}], state, Path(d))
                self.assertEqual([c.args[2] for c in calls.call_args_list], ['prepare'])
            self.assertEqual(state['getadsonline.com']['status'], 'needs_attention')

    def test_unmapped_site_is_reported_while_the_next_supported_site_runs(self):
        from uivision_publish import run_queue
        with tempfile.TemporaryDirectory() as d:
            state = {}
            entries = [{'row': 1, 'domain': 'unmapped.example'}, {'row': 2, 'domain': 'getadsonline.com'}]
            def stage(client, job, action):
                return {'status': 'ready', 'challenge': False} if action == 'prepare' else {'status': 'pending'}
            with patch('uivision_publish.run_stage', side_effect=stage), patch('uivision_publish.prompt') as prompt:
                run_queue(None, None, BUSINESS, entries, state, Path(d), all_sites=True)
                prompt.assert_not_called()
            self.assertEqual(state['unmapped.example']['status'], 'needs_mapping')
            self.assertEqual(state['getadsonline.com']['status'], 'pending')

    def test_human_skip_before_submission_remains_skipped_on_resume(self):
        from uivision_publish import run_queue
        with tempfile.TemporaryDirectory() as d:
            state = {}
            entries = [{'row': 1, 'domain': 'getadsonline.com'}]
            with patch('uivision_publish.run_stage', return_value={'status': 'ready', 'challenge': True}), \
                 patch('uivision_publish.prompt', return_value='skip'):
                run_queue(None, None, BUSINESS, entries, state, Path(d))
            self.assertEqual(state['getadsonline.com']['status'], 'skipped')
            with patch('uivision_publish.run_stage') as stage:
                run_queue(None, None, BUSINESS, entries, state, Path(d))
                stage.assert_not_called()

    def test_operator_defer_preserves_uncertain_submission_and_blocks_future_clicks(self):
        import backlinks
        from uivision_publish import run_queue
        with tempfile.TemporaryDirectory() as d, patch.object(backlinks, 'ROOT', Path(d)):
            workbook, pool = fixtures(Path(d))
            client = inspect_client(workbook, pool)
            folder = client['folder']
            folder.mkdir(parents=True)
            record = {'status': 'uncertain', 'submission_id': 'original', 'reason': 'No receipt'}
            (folder / 'state.json').write_text(json.dumps({'rentry.co': record}))
            self.assertTrue(resolve_review(workbook, pool, 'rentry.co', 'skip'))
            state = json.loads((folder / 'state.json').read_text())
            self.assertEqual(state['rentry.co']['status'], 'uncertain')
            self.assertEqual(state['rentry.co']['submission_id'], 'original')
            self.assertEqual(state['rentry.co']['review_disposition'], 'operator_deferred')
            with patch('uivision_publish.run_stage') as stage, patch('uivision_publish.publish') as individual:
                run_queue(None, None, BUSINESS, [{'domain': 'rentry.co', 'row': 1}], state, folder)
                stage.assert_not_called()
                individual.assert_not_called()

    def test_wrong_hostname_cannot_be_accepted_as_a_manual_result(self):
        import backlinks
        with tempfile.TemporaryDirectory() as d, patch.object(backlinks, 'ROOT', Path(d)):
            workbook, pool = fixtures(Path(d))
            folder = inspect_client(workbook, pool)['folder']
            folder.mkdir(parents=True)
            state = {'rentry.co': {'status': 'uncertain', 'submission_id': 'original'}}
            (folder / 'state.json').write_text(json.dumps(state))
            with self.assertRaises((ValueError, NeedsAttention)):
                resolve_review(workbook, pool, 'rentry.co', 'https://unrelated.test/public')
            self.assertEqual(json.loads((folder / 'state.json').read_text()), state)

    def test_resume_after_review_does_not_repeat_earlier_attempted_sites(self):
        import backlinks
        import uivision_publish
        args = SimpleNamespace(workbook=Path('fixture.xlsx'), pool=Path('pool.xlsx'))
        entries = [{'domain': host} for host in ('first.test', 'review.test', 'next.test')]
        channel = Mock()
        channel.ask.return_value = 'skip'
        calls = []
        def run():
            import sys
            calls.append(list(sys.argv))
            if len(calls) == 1:
                backlinks.begin_attempt(Path('.'), 'first.test')
                raise InspectionRequired('review.test', 'Needs review')
        with patch('desktop_runner.resolve_review', return_value=True), \
             patch('backlinks.load_workbook', return_value=(BUSINESS, entries)), \
             patch('website_pool.merge_pool', return_value=entries), \
             patch.object(backlinks, 'begin_attempt'), patch.object(uivision_publish, 'begin_attempt'), \
             patch.object(uivision_publish, 'main', side_effect=run), patch('sys.argv', ['fixture']):
            publish_from_app(args, channel)
        self.assertEqual(calls[1][calls[1].index('--only') + 1:], ['next.test'])
        self.assertTrue(any(c.args[0] == 'finished' for c in channel.emit.call_args_list))

    def test_native_window_handles_action_and_pause_without_launching_a_browser(self):
        import backlinks
        try:
            root = tk.Tk()
        except tk.TclError:
            self.skipTest('Tk display unavailable')
        root.withdraw()
        try:
            with tempfile.TemporaryDirectory() as d, patch.object(backlinks, 'ROOT', Path(d)):
                workbook, pool = fixtures(Path(d))
                with patch('backlink_app.DEFAULT_WORKBOOK', workbook), patch('backlink_app.DEFAULT_POOL', pool):
                    app = BacklinkApp(root)
                self.assertIsNotNone(app.client)
                self.assertEqual(len(app.table.get_children()), 2)
                app.worker = Mock()
                app.events.put({'event': 'action', 'id': 'request-1', 'message': 'Inspect the result', 'kind': 'review'})
                app.tick()
                root.update_idletasks()
                self.assertEqual(app.action_id, 'request-1')
                self.assertEqual(app.action_frame.winfo_manager(), 'pack')
                app.respond('skip')
                app.worker.send.assert_called_with('respond', id='request-1', value='skip')
                app.pause()
                app.worker.send.assert_called_with('stop')
        finally:
            root.destroy()


if __name__ == '__main__':
    unittest.main()
